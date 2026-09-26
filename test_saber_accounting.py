import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from database import Database
from importer import read_invoices
from lebanese_accounts import LEBANESE_ACCOUNTS
from report_export import export_excel, export_invoice_pdf, export_pdf
from desktop import formatted_user_date, natural_sort_value, parse_user_date, row_matches_search, sortable_date
from company_manager import CompanyManager
from run_desktop import local_server_ready

class SaberAccountingTest(unittest.TestCase):
    def test_local_server_probe_is_safe_when_service_is_not_running(self):
        self.assertIn(local_server_ready(),(True,False))

    def test_sort_helpers_support_dates_numbers_and_names(self):
        dates=["31-12-2025","01-01-2026","15-06-2024"]
        self.assertEqual(sorted(dates,key=sortable_date),["15-06-2024","31-12-2025","01-01-2026"])
        numbers=["JV-10","2","11","1"]
        self.assertEqual(sorted(numbers,key=natural_sort_value)[:3],["1","2","11"])
        self.assertEqual(formatted_user_date("31122024"),"31-12-2024")
        self.assertEqual(parse_user_date("01012025").year,2025)

    def test_table_search_matches_all_terms_across_columns(self):
        row=("INV-100","22-09-2026","Supplier Alpha","purchase","USD",100,11,111)
        self.assertTrue(row_matches_search(row,"supplier usd"))
        self.assertTrue(row_matches_search(row,"INV-100 111"))
        self.assertTrue(row_matches_search(row,""))
        self.assertFalse(row_matches_search(row,"supplier eur"))
        self.assertFalse(row_matches_search(row,"missing"))

    def test_import_rules_and_balancing(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "sample.xlsx"
            wb = Workbook(); ws = wb.active
            ws.append(["Invoice Number","Date","Supplier Name","Total Before VAT","VAT","Total After VAT"])
            ws.append(["A-100","01-09-2026","Supplier A",100,11,111])
            ws.append([None,None,None,None,None,None])
            ws.append([None,"02-09-2026","Supplier B",200,22,200])
            ws.append(["A-100","01-09-2026","Supplier A",100,11,111]); wb.save(path)
            rows = read_invoices(path)
            self.assertEqual([r["invoice_number"] for r in rows], ["A-100","4","A-100"])
            db=Database(Path(folder)/"test.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            for row in rows: db.import_invoice(row,user["id"])
            self.assertEqual(len(db.list_invoices()),3)
            tb=db.trial_balance()
            self.assertAlmostEqual(sum(float(r["debit"] or 0) for r in tb),sum(float(r["credit"] or 0) for r in tb),places=2)

    def test_currency_detection_defaults_conflicts_and_separation(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "currencies.xlsx"
            wb = Workbook(); ws = wb.active
            ws.append(["Invoice Number","Date","Supplier Name","Total Before VAT","VAT","Total After VAT","Currency"])
            ws.append(["USD-1","01-09-2026","A","$100","$11","$111",None])
            ws.append(["EUR-1","02-09-2026","B","€200","€22","€222",None])
            ws.append(["LBP-1","03-09-2026","C","300 L.L.","33 L.L.","333 L.L.",None])
            ws.append(["AED-1","04-09-2026","D","AED 400","AED 44","AED 444",None])
            ws.append(["DEFAULT-1","05-09-2026","E",500,55,555,None])
            ws.append(["CONFLICT-1","06-09-2026","F","$600","$66","$666","EUR"])
            ws.append(["PRICE-WINS","07-09-2026","G","€700","€77","€777","USD"])
            wb.save(path)
            rows = read_invoices(path)
            self.assertEqual([r["currency"] for r in rows[:5]], ["USD","EUR","LBP","AED","USD"])
            self.assertEqual(rows[6]["currency"], "EUR")
            self.assertEqual(rows[6]["currency_issue"], "")
            self.assertEqual(rows[4]["currency_issue"], "missing_defaulted_to_usd")
            self.assertEqual(rows[5]["currency_issue"], "")

            db=Database(Path(folder)/"currency.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            for row in rows: db.import_invoice(row,user["id"])
            self.assertEqual({r["currency"] for r in db.dashboard()}, {"USD","EUR","LBP","AED"})
            trial=db.trial_balance()
            self.assertEqual({r["currency"] for r in trial}, {"USD","EUR","LBP","AED"})
            conflict=next(r for r in db.list_invoices() if r["invoice_number"]=="CONFLICT-1")
            self.assertEqual(conflict["status"], "posted")

    def test_account_numbers_import_and_journal_posting(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "accounts.xlsx"
            wb = Workbook(); ws = wb.active
            ws.append([
                "Invoice Number", "Date", "Supplier Name",
                "Total Before VAT", "VAT", "Total After VAT",
                "Supplier Account Number", "VAT Account Number",
                "Expense Account Number",
            ])
            ws.append([
                "ACC-1", "15-09-2026", "Supplier", "$100", "$11", "$111",
                "2110", "1310", "5110",
            ])
            wb.save(path)
            invoice = read_invoices(path)[0]
            self.assertEqual(invoice["supplier_account"], "2110")
            self.assertEqual(invoice["vat_account"], "1310")
            self.assertEqual(invoice["expense_account"], "5110")

            db = Database(Path(folder) / "accounts.db")
            db.initialize("secret")
            user = db.user_for_token(db.login("admin", "secret")["token"])
            db.import_invoice(invoice, user["id"])
            saved = db.list_invoices()[0]
            self.assertEqual(saved["supplier_account"], "2110")
            self.assertEqual(saved["vat_account"], "1310")
            self.assertEqual(saved["expense_account"], "5110")
            trial_codes = {row["code"] for row in db.trial_balance()}
            self.assertTrue({"2110", "1310", "5110"}.issubset(trial_codes))

    def test_debit_credit_visible_invoices_dashboard_and_journal(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"journal.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            db.import_invoice({"invoice_number":"S-1","invoice_date":"20-09-2026","party_name":"Customer",
                "kind":"sale","currency":"USD","subtotal":100,"vat":11,"total":111},user["id"])
            db.import_invoice({"invoice_number":"P-1","invoice_date":"21-09-2026","party_name":"Supplier",
                "kind":"purchase","currency":"USD","subtotal":200,"vat":22,"total":222},user["id"])
            invoices={row["invoice_number"]:row for row in db.list_invoices()}
            self.assertEqual(invoices["S-1"]["debit"],111.0)
            self.assertEqual(invoices["S-1"]["credit"],0)
            self.assertEqual(invoices["P-1"]["debit"],0)
            self.assertEqual(invoices["P-1"]["credit"],222.0)
            dashboard={row["kind"]:row for row in db.dashboard()}
            self.assertEqual(dashboard["sale"]["debit"],111.0)
            self.assertEqual(dashboard["purchase"]["credit"],222.0)
            journal=db.journal("2026-09-20","2026-09-21","USD")
            self.assertEqual(len(journal),6)
            self.assertAlmostEqual(sum(row["debit"] for row in journal),333.0)
            self.assertAlmostEqual(sum(row["credit"] for row in journal),333.0)
            self.assertTrue(all("balance" in row for row in journal))

    def test_professional_invoice_lifecycle(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"lifecycle.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            invoice={"invoice_number":"","invoice_date":"22-09-2026","party_name":"Customer A","kind":"sale","currency":"USD"}
            invoice_id=db.create_manual_invoice(invoice,[{"description":"Service","quantity":1,"unit_price":100,"vat_rate":11}],user["id"])
            saved=db.get_invoice(invoice_id)
            self.assertEqual(saved["invoice_number"],"SAL-2026-000001")
            self.assertEqual(saved["payment_status"],"unpaid")
            updated=db.update_invoice(invoice_id,{**saved,"amount_paid":50,"status":"posted"},user["id"])
            self.assertEqual(updated["payment_status"],"partial")
            self.assertEqual(updated["outstanding"],61.0)
            duplicate=db.duplicate_invoice(invoice_id,user["id"])
            self.assertEqual(duplicate["invoice_number"],"SAL-2026-000002")
            cancelled=db.cancel_invoice(invoice_id,"Customer request",user["id"])
            self.assertEqual(cancelled["status"],"cancelled")
            journal=db.journal(currency="USD")
            debit=sum(row["debit"] for row in journal); credit=sum(row["credit"] for row in journal)
            self.assertAlmostEqual(debit,credit)
            self.assertTrue(any(row["entry_number"].startswith("REV-") for row in journal))
            attachment_id=db.add_attachment(duplicate["id"],"invoice.pdf","application/pdf",b"PDF",user["id"])
            self.assertEqual(db.get_attachment(attachment_id)["content"],b"PDF")
            self.assertGreaterEqual(len(db.invoice_history(invoice_id)),2)

    def test_profit_loss_and_close_fiscal_year(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"year.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            db.import_invoice({"invoice_number":"S-1","invoice_date":"15-06-2026","party_name":"Client","kind":"sale","currency":"USD","subtotal":1000,"vat":110,"total":1110},user["id"])
            db.import_invoice({"invoice_number":"P-1","invoice_date":"20-06-2026","party_name":"Supplier","kind":"purchase","currency":"USD","subtotal":400,"vat":44,"total":444},user["id"])
            pnl=db.profit_and_loss("2026-01-01","2026-12-31","USD")
            income=sum(row["amount"] for row in pnl if row["type"]=="income")
            expenses=sum(row["amount"] for row in pnl if row["type"]=="expense")
            self.assertEqual(income-expenses,600.0)
            result=db.close_fiscal_year(2026,user["id"])
            self.assertEqual(result["opened_year"],2027)
            self.assertEqual(result["net_results"]["USD"],600.0)
            self.assertEqual(db.profit_and_loss("2027-01-01","2027-12-31","USD"),[])
            years={row["year"]:row["status"] for row in db.list_fiscal_years()}
            self.assertEqual(years[2026],"closed"); self.assertEqual(years[2027],"open")

    def test_payments_expenses_reports_and_period_lock(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"professional.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            customer=db.save_party({"name":"Customer A","kind":"customer","tax_number":"C-1","currency":"USD"},user["id"])
            supplier=db.save_party({"name":"Supplier A","kind":"supplier","tax_number":"S-1","currency":"USD"},user["id"])
            db.import_invoice({"invoice_number":"S-1","invoice_date":"01-06-2026","party_name":"Customer A","kind":"sale","currency":"USD","subtotal":1000,"vat":110,"total":1110},user["id"])
            db.import_invoice({"invoice_number":"P-1","invoice_date":"02-06-2026","party_name":"Supplier A","kind":"purchase","currency":"USD","subtotal":400,"vat":44,"total":444},user["id"])
            db.add_payment({"kind":"customer_receipt","party_id":customer["id"],"payment_date":"03-06-2026","currency":"USD","amount":500,"cash_account":"531","party_account":"4111"},user["id"])
            db.add_payment({"kind":"supplier_payment","party_id":supplier["id"],"payment_date":"04-06-2026","currency":"USD","amount":200,"cash_account":"5121","party_account":"4011"},user["id"])
            db.add_expense({"expense_date":"05-06-2026","description":"Office expense","category":"Office","currency":"USD","subtotal":100,"vat":11,"expense_account":"6011","vat_account":"4426.6","payment_account":"531"},user["id"])
            self.assertEqual(len(db.list_payments()),2); self.assertEqual(len(db.list_expenses()),1)
            ledger=db.general_ledger("531","2026-01-01","2026-12-31","USD")
            self.assertTrue(ledger["items"])
            balance=db.balance_sheet("2026-12-31","USD")
            self.assertTrue(any(row["code"]=="13" for row in balance))
            vat=db.vat_report("2026-01-01","2026-12-31","USD")["summary"][0]
            self.assertEqual(vat["sales_vat"],110.0); self.assertEqual(vat["recoverable_vat"],55.0); self.assertEqual(vat["vat_payable"],55.0)
            db.close_fiscal_year(2026,user["id"])
            with self.assertRaisesRegex(ValueError,"closed"):
                db.add_expense({"expense_date":"31-12-2026","description":"Late","currency":"USD","subtotal":1,"vat":0},user["id"])

    def test_security_backup_rates_dashboard_and_branded_invoice(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"complete.db"); db.initialize("secret")
            admin=db.user_for_token(db.login("admin","secret")["token"])
            saved=db.save_user({"username":"viewer1","password":"Viewer123!","role":"viewer","language":"fr","active":True},admin["id"])
            self.assertEqual(saved["role"],"viewer"); self.assertEqual(len(db.list_users()),2)
            db.save_settings({"base_currency":"EUR","backup_interval_hours":"12"},admin["id"])
            self.assertEqual(db.settings()["base_currency"],"EUR")
            db.save_exchange_rate({"rate_date":"22-09-2026","from_currency":"USD","to_currency":"LBP","rate":"89500"},admin["id"])
            rates=db.list_exchange_rates(); fixed=next(row for row in rates if row["from_currency"]=="USD" and row["to_currency"]=="LBP")
            self.assertEqual(fixed["rate"],89500.0)
            result=db.save_exchange_rate({"date_from":"01-01-2025","date_to":"03-01-2025","from_currency":"EUR","to_currency":"USD","rate":"1.04"},admin["id"])
            self.assertEqual(result["days"],3)
            backup=db.backup(); self.assertTrue(Path(backup).exists()); self.assertTrue(db.list_backups())
            invoice_id=db.create_manual_invoice({"invoice_number":"","invoice_date":"22-09-2026","party_name":"Client","kind":"sale","currency":"USD","due_date":"01-09-2026"},
                [{"description":"Audit service","quantity":1,"unit_price":100,"vat_rate":11}],admin["id"])
            dashboard=db.professional_dashboard(); self.assertEqual(dashboard["metrics"][0]["sales"],100.0); self.assertEqual(dashboard["metrics"][0]["overdue"],1)
            detail=db.invoice_detail(invoice_id); pdf=Path(folder)/"invoice.pdf"
            export_invoice_pdf(pdf,detail["invoice"],detail["items"])
            self.assertGreater(pdf.stat().st_size,1000)

    def test_manual_invoice_items_editable_subtotal_and_vat(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"manual.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            invoice={"invoice_number":"M-1","invoice_date":"14-09-2026","party_name":"Manual Supplier",
                     "kind":"purchase","currency":"USD","source_file":"Journal Voucher"}
            items=[
                {"description":"Automatic VAT","quantity":2,"unit_price":50,"vat_rate":11},
                {"description":"Edited subtotal","quantity":1,"unit_price":100,"subtotal":90,"vat_rate":5},
                {"description":"Edited VAT","quantity":1,"unit_price":20,"vat_rate":11,"vat":1},
            ]
            invoice_id=db.create_manual_invoice(invoice,items,user["id"])
            saved=next(row for row in db.list_invoices() if row["id"]==invoice_id)
            self.assertEqual(float(saved["subtotal"]),210.0)
            self.assertEqual(float(saved["vat"]),16.5)
            self.assertEqual(float(saved["total"]),226.5)
            with db.connect() as connection:
                lines=connection.execute("SELECT * FROM invoice_items WHERE invoice_id=? ORDER BY id",(invoice_id,)).fetchall()
            self.assertEqual(len(lines),3)
            self.assertEqual(float(lines[1]["subtotal"]),90.0)
            self.assertEqual(float(lines[2]["vat"]),1.0)
            voucher=next(row for row in db.journal() if row["source_id"]==invoice_id)
            db.delete_journal_voucher(voucher["entry_id"],user["id"])
            self.assertFalse(any(row["id"]==invoice_id for row in db.list_invoices()))

    def test_sales_invoice_draft_and_custom_nine_digit_posting_accounts(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"sales_invoice.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            invoice={"invoice_number":"","invoice_date":"24-09-2026","party_name":"Sales Client","kind":"sales","currency":"USD",
                     "supplier_account":"","vat_account":"442700000","expense_account":"713100000","status":"review","source_file":"Sales Invoice"}
            invoice_id=db.create_manual_invoice(invoice,[{"description":"Professional service","quantity":2,"unit_price":100,"vat_rate":11}],user["id"])
            row=next(item for item in db.list_invoices() if item["id"]==invoice_id)
            self.assertEqual(row["status"],"review")
            self.assertTrue(str(row["invoice_number"]).startswith("SAL-2026-"))
            self.assertEqual(len(str(row["supplier_account"])),9)
            codes={item["code"] for item in db.trial_balance(posting_status="review")}
            self.assertTrue({str(row["supplier_account"]),"442700000","713100000"}.issubset(codes))

    def test_legal_documents_and_three_document_case_workflows(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"documents.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            supplier=db.save_party({"kind":"supplier","name":"Customs Supplier","currency":"USD"},user["id"])
            document_id=db.add_party_document(supplier["id"],{"document_type":"MOF / VAT Certificate","issue_date":"01-01-2026","expiry_date":"31-12-2026","file_name":"mof.pdf","mime_type":"application/pdf"},b"PDF",user["id"])
            self.assertEqual(db.get_party_document(document_id)["content"],b"PDF")
            for case_type,roles in (("purchase",("supplier_invoice",)),("expense",("expense_document",)),("customs",("supplier_invoice","customs_declaration","broker_invoice"))):
                case=db.save_document_case({"case_type":case_type,"document_date":"24-09-2026","party_id":supplier["id"],"currency":"USD",
                    "supplier_invoice_amount":100,"freight":10 if case_type=="customs" else 0,"insurance":5 if case_type=="customs" else 0,
                    "customs_duties":20 if case_type=="customs" else 0,"import_vat":11,"broker_fees":5 if case_type=="customs" else 0},user["id"])
                self.assertEqual(case["status"],"draft")
                with self.assertRaisesRegex(ValueError,"Attach required"):
                    db.post_document_case(case["id"],user["id"])
                for role in roles: db.add_case_attachment(case["id"],role,role+".pdf","application/pdf",b"PDF",user["id"])
                posted=db.post_document_case(case["id"],user["id"])
                self.assertEqual(posted["status"],"posted"); self.assertIsNotNone(posted["invoice_id"])
                self.assertEqual(len(db.list_attachments(posted["invoice_id"])),len(roles))

    def test_cash_flow_aging_and_comparative_reports(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"advanced_reports.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            customer=db.save_party({"kind":"customer","name":"Report Customer","currency":"USD"},user["id"])
            db.import_invoice({"invoice_number":"AR-1","invoice_date":"01-01-2026","due_date":"15-01-2026","party_name":"Report Customer","kind":"sale","currency":"USD","subtotal":100,"vat":11,"total":111,"status":"posted"},user["id"])
            db.import_invoice({"invoice_number":"AR-0","invoice_date":"01-01-2025","party_name":"Report Customer","kind":"sale","currency":"USD","subtotal":80,"vat":8.8,"total":88.8,"status":"posted"},user["id"])
            db.add_payment({"kind":"customer_receipt","payment_date":"20-01-2026","party_id":customer["id"],"currency":"USD","amount":50,"cash_account":"531","party_account":customer["account_number"]},user["id"])
            aging=db.aging_report("20-02-2026","sale","USD")
            self.assertTrue(any(row["invoice_number"]=="AR-1" and row["bucket"]=="31-60" for row in aging))
            cash=db.cash_flow("2026-01-01","2026-12-31","USD")
            self.assertAlmostEqual(sum(row["net"] for row in cash),50.0)
            comparative=db.comparative_reports("2026-01-01","2026-12-31","USD")
            sales=next(row for row in comparative["items"] if row["type"]=="income")
            self.assertAlmostEqual(sales["current"],100.0); self.assertAlmostEqual(sales["prior"],80.0)

    def test_direct_journal_voucher_new_edit_multiple_lines_and_number(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"voucher.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            saved=db.save_journal_voucher({"entry_date":"22-09-2026","description":"Opening adjustment","currency":"USD"},[
                {"account_code":"531","description":"Cash","debit":100,"credit":0},
                {"account_code":"601100000","description":"Expense reversal","debit":50,"credit":0},
                {"account_code":"4011","description":"Supplier","debit":0,"credit":150}],user["id"])
            self.assertEqual(saved["voucher"]["entry_number"],"JV-2026-000001"); self.assertEqual(len(saved["lines"]),3)
            entry_id=saved["voucher"]["id"]
            edited=db.save_journal_voucher({"entry_number":"JV-CUSTOM-1","entry_date":"22-09-2026","description":"Edited adjustment","currency":"USD"},[
                {"account_code":"531","debit":200,"credit":0},{"account_code":"4011","debit":0,"credit":200}],user["id"],entry_id)
            self.assertEqual(edited["voucher"]["entry_number"],"JV-CUSTOM-1"); self.assertEqual(len(edited["lines"]),2)
            db.delete_journal_voucher(entry_id,user["id"])
            with self.assertRaises(KeyError): db.journal_voucher_detail(entry_id)

    def test_branches_filter_invoices_vouchers_trial_and_statement(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"branches.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            north=db.save_branch({"name":"North Branch"},user["id"]); south=db.save_branch({"name":"South Branch"},user["id"])
            invoice_id=db.import_invoice({"invoice_number":"BR-1","invoice_date":"22-09-2026","party_name":"Branch Supplier","kind":"purchase","currency":"USD","subtotal":100,"vat":11,"total":111,"branch_id":north["id"]},user["id"])
            invoice=next(row for row in db.list_invoices() if row["id"]==invoice_id)
            self.assertEqual(invoice["branch_name"],"North Branch")
            voucher=db.save_journal_voucher({"entry_date":"22-09-2026","description":"South adjustment","currency":"USD","branch_id":south["id"]},[{"account_code":"531","debit":10,"credit":0},{"account_code":"4011","debit":0,"credit":10}],user["id"])
            self.assertTrue(db.trial_balance(branch_id=north["id"])); self.assertTrue(db.trial_balance(branch_id=south["id"]))
            self.assertTrue(all(row["branch_id"]==south["id"] for row in db.journal(branch_id=south["id"])))
            party=next(row for row in db.list_parties() if row["name"]=="Branch Supplier")
            self.assertEqual(len(db.statement_of_account(party["id"],branch_id=north["id"])["items"]),1)
            self.assertEqual(len(db.statement_of_account(party["id"],branch_id=south["id"])["items"]),0)

    def test_party_account_categories_and_prefix_sequence(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"party-types.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            client=db.save_party({"name":"Client Category","account_category":"client","currency":"USD","account_number":"4119"},user["id"])
            asset=db.save_party({"name":"Asset Supplier Category","account_category":"asset_supplier","currency":"USD","account_number":"4019"},user["id"])
            payable=db.save_party({"name":"Other Payable Category","account_category":"other_payable","currency":"USD","account_number":"4019"},user["id"])
            self.assertEqual(client["account_number"],"411900001")
            self.assertEqual(asset["account_number"],"401900001"); self.assertEqual(payable["account_number"],"401900002")
            categories={row["name"]:row["account_category"] for row in db.list_parties()}
            self.assertEqual(categories["Client Category"],"client"); self.assertEqual(categories["Asset Supplier Category"],"asset_supplier")
            auto_asset=db.save_party({"name":"Automatic Asset Supplier","account_category":"asset_supplier","currency":"USD"},user["id"])
            auto_payable=db.save_party({"name":"Automatic Other Payable","account_category":"other_payable","currency":"USD"},user["id"])
            self.assertTrue(auto_asset["account_number"].startswith("4031")); self.assertTrue(auto_payable["account_number"].startswith("4619"))

    def test_currency_detection_from_excel_number_formats(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "formatted-currencies.xlsx"
            wb = Workbook(); ws = wb.active
            ws.append(["Invoice Number","Date","Supplier Name","Total Before VAT","VAT","Total After VAT"])
            formats = [
                ("USD-FMT", '"$"#,##0.00'),
                ("EUR-FMT", '€#,##0.00'),
                ("LBP-FMT", '#,##0.00 "L.L."'),
                ("AED-FMT", '#,##0.00 "AED"'),
            ]
            for invoice_number, number_format in formats:
                ws.append([invoice_number,"14-09-2026","Supplier",100,11,111])
                for column in (4,5,6):
                    ws.cell(ws.max_row,column).number_format=number_format
            ws.append(["EUR-BLANK-VAT","14-09-2026","Supplier",100,None,100])
            ws.cell(ws.max_row,4).number_format='€#,##0.00'
            ws.cell(ws.max_row,5).number_format='"$"#,##0.00'
            ws.cell(ws.max_row,6).number_format='€#,##0.00'
            ws.append(["EUR-MAJORITY","14-09-2026","Supplier",100,11,111])
            ws.cell(ws.max_row,4).number_format='€#,##0.00'
            ws.cell(ws.max_row,5).number_format='"$"#,##0.00'
            ws.cell(ws.max_row,6).number_format='€#,##0.00'
            wb.save(path)
            rows=read_invoices(path)
            self.assertEqual([row["currency"] for row in rows],["USD","EUR","LBP","AED","EUR","EUR"])
            self.assertNotIn("conflicting",rows[4]["currency_issue"])
            self.assertEqual(rows[5]["currency_issue"], "")

    def test_add_item_and_statement_of_account(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"statement.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            first={"invoice_number":"P-1","invoice_date":"01-09-2026","party_name":"Supplier A",
                   "kind":"purchase","currency":"USD","subtotal":100,"vat":11,"total":111}
            second={"invoice_number":"P-2","invoice_date":"15-09-2026","party_name":"Supplier A",
                    "kind":"purchase","currency":"USD","subtotal":200,"vat":22,"total":222}
            first_id=db.import_invoice(first,user["id"]); db.import_invoice(second,user["id"])
            updated=db.add_invoice_item(first_id,{"description":"Extra","quantity":1,"unit_price":50,
                "subtotal":50,"vat_rate":10,"vat":5},user["id"])
            self.assertEqual(float(updated["total"]),166.0)
            party=next(row for row in db.list_parties() if row["name"]=="Supplier A")
            statement=db.statement_of_account(party["id"],"2026-09-10","2026-09-30","USD")
            self.assertEqual(statement["opening"]["USD"],-166.0)
            self.assertEqual(len(statement["items"]),1)
            self.assertEqual(statement["items"][0]["credit"],222.0)
            self.assertEqual(statement["items"][0]["balance"],-388.0)
            trial=db.trial_balance()
            self.assertAlmostEqual(sum(float(row["debit"] or 0) for row in trial),
                                   sum(float(row["credit"] or 0) for row in trial),places=2)

    def test_full_lebanese_chart_and_default_posting(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"lebanese.db"); db.initialize("secret")
            accounts=db.list_accounts()
            self.assertEqual(len(accounts),len(LEBANESE_ACCOUNTS)+6+19)  # includes sale/credit-note posting + bank commission & exchange-difference accounts
            codes={row["code"] for row in accounts}
            self.assertTrue({"1","2","3","4","5","6","7","4011","4111","4426.6","4427","6011","713",
                             "7011","701100001","711100001","7130","713000001","709000001","7190","719000001"}.issubset(codes))
            self.assertEqual(next(row["parent_code"] for row in accounts if row["code"]=="719000001"),"7190")
            self.assertFalse({"1100","2100","2200","1300","4100","5100","9999"} & codes)
            user=db.user_for_token(db.login("admin","secret")["token"])
            purchase={"invoice_number":"LB-P","invoice_date":"21-09-2026","party_name":"Supplier",
                      "kind":"purchase","currency":"USD","subtotal":100,"vat":11,"total":111}
            sale={"invoice_number":"LB-S","invoice_date":"21-09-2026","party_name":"Customer",
                  "kind":"sale","currency":"USD","subtotal":200,"vat":22,"total":222}
            db.import_invoice(purchase,user["id"]); db.import_invoice(sale,user["id"])
            trial_codes={row["code"] for row in db.trial_balance()}
            self.assertTrue({"4111","44210","4427","601100000","713"}.issubset(trial_codes))
            self.assertTrue(any(code.startswith("4011") and len(code)==9 for code in trial_codes))

    def test_suppliers_receive_unique_nine_digit_accounts(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"suppliers.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            first=db.save_party({"name":"Supplier One","kind":"supplier","currency":"USD"},user["id"])
            second=db.save_party({"name":"Supplier Two","kind":"supplier","currency":"EUR"},user["id"])
            customer=db.save_party({"name":"Customer One","kind":"customer","currency":"USD"},user["id"])
            prefixed_one=db.save_party({"name":"Other Payable One","kind":"supplier","currency":"USD","account_number":"4999"},user["id"])
            prefixed_two=db.save_party({"name":"Other Payable Two","kind":"supplier","currency":"USD","account_number":"4999"},user["id"])
            self.assertRegex(first["account_number"],r"^\d{9}$")
            self.assertRegex(second["account_number"],r"^\d{9}$")
            self.assertNotEqual(first["account_number"],second["account_number"])
            self.assertRegex(customer["account_number"],r"^4111\d{5}$")
            self.assertEqual(prefixed_one["account_number"],"499900001")
            self.assertEqual(prefixed_two["account_number"],"499900002")
            invoice_id=db.import_invoice({"invoice_number":"AUTO-AC","invoice_date":"22-09-2026",
                "party_name":"Supplier One","kind":"purchase","currency":"USD","subtotal":100,"vat":11,"total":111},user["id"])
            invoice=next(row for row in db.list_invoices() if row["id"]==invoice_id)
            self.assertEqual(invoice["supplier_account"],first["account_number"])

    def test_new_types_no_vat_editable_dc_and_nine_digit_accounts(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"new_features.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            db.save_account({"code":"601100123","name_en":"Consulting expense","type":"expense","parent_code":"6011"},user["id"])
            self.assertIn("601100123",{row["code"] for row in db.list_accounts()})
            invoice_id=db.import_invoice({"invoice_number":"NV-1","invoice_date":"22-09-2026","party_name":"Supplier NV",
                "kind":"expenses","entry_type":"expenses","currency":"USD","subtotal":100,"vat":0,"total":100},user["id"])
            saved=next(row for row in db.list_invoices() if row["id"]==invoice_id)
            self.assertEqual(saved["vat"],"0"); self.assertEqual(saved["total"],"100")
            self.assertEqual(saved["expense_account"],"601100000"); self.assertEqual(saved["expense_no_vat_account"],"601100001"); self.assertEqual(saved["vat_account"],"44216")  # expenses VAT account
            changed=db.update_invoice(invoice_id,{"invoice_number":"NV-1","invoice_date":"22-09-2026","party_name":"Supplier NV",
                "kind":"assets","entry_type":"assets","currency":"USD","subtotal":100,"vat":0,"total":100,"debit":75,"credit":25,
                "supplier_account":saved["supplier_account"],"vat_account":saved["vat_account"],"expense_account":"601100123","expense_no_vat_account":"601100001",
                "supplier_side":"D - Debit","vat_side":"C - Credit","expense_side":"C - Credit","status":"posted"},user["id"])
            listed=next(row for row in db.list_invoices() if row["id"]==invoice_id)
            self.assertEqual(changed["entry_type"],"assets"); self.assertEqual(listed["debit"],75); self.assertEqual(listed["credit"],25)
            self.assertEqual(listed["supplier_side"],"D"); self.assertEqual(listed["vat_side"],"C"); self.assertEqual(listed["expense_side"],"C")

    def test_deductible_split_and_automatic_unique_account_number(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"deductible.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            created=db.save_account({"code":"","name_en":"Auto expense","type":"expense","parent_code":"6011"},user["id"])
            self.assertEqual(created["code"],"601100002")
            self.assertEqual(db.next_account_number("6011"),"601100003")
            prefixed=db.save_account({"code":"6011","name_en":"Auto from prefix","type":"expense"},user["id"])
            self.assertEqual(prefixed["code"],"601100003")
            with self.assertRaisesRegex(ValueError,"already exists"):
                db.save_account({"code":"601100002","name_en":"Duplicate","type":"expense","parent_code":"6011"},user["id"])
            invoice_id=db.create_manual_invoice({"invoice_number":"SPLIT-1","invoice_date":"22-09-2026","party_name":"Split Supplier","kind":"expenses","currency":"USD"},[
                {"description":"Mixed expense","quantity":1,"unit_price":80,"deductible_subtotal":80,"non_deductible_subtotal":20,"vat_rate":11}
            ],user["id"])
            invoice=next(row for row in db.list_invoices() if row["id"]==invoice_id)
            self.assertEqual(float(invoice["deductible_subtotal"]),80); self.assertEqual(float(invoice["non_deductible_subtotal"]),20)
            self.assertEqual(float(invoice["vat"]),8.8); self.assertEqual(float(invoice["total"]),108.8)
            journal=db.journal(currency="USD")
            by_account={row["account_code"]:row for row in journal}
            self.assertEqual(by_account["601100000"]["debit"],80); self.assertEqual(by_account["601100001"]["debit"],20)

    def test_account_name_edit_and_split_expense(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"expense_split.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            created=db.save_account({"code":"","name_en":"Original name","type":"expense","parent_code":"6011"},user["id"])
            db.rename_account(created["code"],"Updated expense name",user["id"])
            account=next(row for row in db.list_accounts() if row["code"]==created["code"])
            self.assertEqual(account["name_en"],"Updated expense name")
            expense_id=db.add_expense({"expense_date":"22-09-2026","description":"Mixed office expense","category":"Office","currency":"EUR",
                "with_vat_subtotal":80,"without_vat_subtotal":20,"vat":8.8,"expense_account":"601100000","expense_without_vat_account":"601100001",
                "vat_account":"442660000","payment_account":"531"},user["id"])
            saved=next(row for row in db.list_expenses() if row["id"]==expense_id)
            self.assertEqual(saved["with_vat_subtotal"],80); self.assertEqual(saved["without_vat_subtotal"],20); self.assertEqual(saved["total"],108.8)
            journal=[row for row in db.journal(currency="EUR") if row["source_type"]=="expense"]
            debits={row["account_code"]:row["debit"] for row in journal}
            self.assertEqual(debits["601100000"],80); self.assertEqual(debits["601100001"],20); self.assertEqual(debits["442660000"],8.8)

    def test_update_specific_invoice_row(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"update.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            invoice={"invoice_number":"OLD-1","invoice_date":"01-09-2026","party_name":"Old Supplier",
                     "kind":"purchase","currency":"USD","subtotal":100,"vat":11,"total":111}
            invoice_id=db.import_invoice(invoice,user["id"])
            updated=db.update_invoice(invoice_id,{
                "invoice_number":"NEW-1","invoice_date":"20-09-2026","party_name":"New Supplier",
                "kind":"purchase","currency":"EUR","subtotal":"200","vat":"22","total":"222",
                "supplier_account":"2110","vat_account":"1310","expense_account":"5110","status":"posted",
            },user["id"])
            self.assertEqual(updated["invoice_number"],"NEW-1")
            self.assertEqual(updated["currency"],"EUR")
            self.assertEqual(updated["party_name"],"New Supplier")
            trial=db.trial_balance("2026-09-20","2026-09-20")
            self.assertEqual({row["currency"] for row in trial},{"EUR"})
            self.assertAlmostEqual(sum(float(row["debit"] or 0) for row in trial),222.0)
            self.assertAlmostEqual(sum(float(row["credit"] or 0) for row in trial),222.0)
            self.assertTrue({"2110","1310","5110"}.issubset({row["code"] for row in trial}))

    def test_trial_balance_date_range(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"dates.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            rows=[
                {"invoice_number":"D-1","invoice_date":"01-09-2026","party_name":"A","kind":"purchase","currency":"USD","subtotal":100,"vat":11,"total":111},
                {"invoice_number":"D-2","invoice_date":"2026-09-15","party_name":"B","kind":"purchase","currency":"USD","subtotal":200,"vat":22,"total":222},
                {"invoice_number":"D-3","invoice_date":"30-09-2026","party_name":"C","kind":"purchase","currency":"USD","subtotal":300,"vat":33,"total":333},
            ]
            for row in rows: db.import_invoice(row,user["id"])
            filtered=db.trial_balance("2026-09-10","2026-09-20")
            self.assertTrue(filtered)
            self.assertAlmostEqual(sum(float(r["debit"] or 0) for r in filtered),222.0)
            self.assertAlmostEqual(sum(float(r["credit"] or 0) for r in filtered),222.0)
            self.assertTrue(any(abs(float(r["opening"]))>0 for r in filtered))
            account=filtered[0]["code"]
            exact=db.trial_balance("2026-09-10","2026-09-20",account,False)
            self.assertTrue(exact); self.assertTrue(all(r["code"]==account for r in exact))
            self.assertTrue(all("usd_opening" in r and "lbp_closing_balance" in r for r in exact))
            ranged=db.trial_balance("2026-09-10","2026-09-20",None,True,"4011","601199999")
            self.assertTrue(ranged); self.assertTrue(all(4011<=int(''.join(c for c in r["code"] if c.isdigit()))<=601199999 for r in ranged))

    def test_trial_balance_posted_review_and_both_filters(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"posting_filter.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            db.import_invoice({"invoice_number":"POST-1","invoice_date":"01-01-2026","party_name":"Posted Client","kind":"sale","currency":"USD","subtotal":100,"vat":11,"total":111,"status":"posted"},user["id"])
            db.import_invoice({"invoice_number":"REV-1","invoice_date":"02-01-2026","party_name":"Review Client","kind":"sale","currency":"USD","subtotal":200,"vat":22,"total":222,"status":"review"},user["id"])
            posted=db.trial_balance(posting_status="posted"); review=db.trial_balance(posting_status="review"); both=db.trial_balance(posting_status="both")
            self.assertTrue(posted); self.assertTrue(review)
            self.assertGreater(sum(abs(r["balance"]) for r in both),sum(abs(r["balance"]) for r in posted))

    def test_report_exports(self):
        with tempfile.TemporaryDirectory() as folder:
            rows=[["purchase","USD",2,300,33,333]]; headers=["Type","Currency","Invoices","Before VAT","VAT","Total"]
            xlsx=Path(folder)/"dashboard.xlsx"; pdf=Path(folder)/"dashboard.pdf"
            export_excel(xlsx,"Dashboard",headers,rows); export_pdf(pdf,"Dashboard",headers,rows)
            self.assertGreater(xlsx.stat().st_size,1000); self.assertGreater(pdf.stat().st_size,500)

    def test_replacement_removes_previous_invoice_set(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"test.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            first={"invoice_number":"1","invoice_date":"2026-09-01","party_name":"A","kind":"purchase","currency":"USD","subtotal":100,"vat":11,"total":111}
            second={"invoice_number":"2","invoice_date":"2026-09-02","party_name":"B","kind":"purchase","currency":"USD","subtotal":200,"vat":22,"total":222}
            db.import_invoice(first,user["id"]); result=db.clear_invoices(user["id"]); db.import_invoice(second,user["id"])
            self.assertEqual(result["deleted"],1); self.assertTrue(Path(result["backup"]).exists())
            invoices=db.list_invoices(); self.assertEqual(len(invoices),1); self.assertEqual(invoices[0]["invoice_number"],"2")
            db.delete_invoice(invoices[0]["id"],user["id"]); self.assertFalse(db.list_invoices())

    def test_multi_company_and_separate_fiscal_year_databases(self):
        with tempfile.TemporaryDirectory() as folder:
            master=Database(Path(folder)/"master.db"); master.initialize("secret")
            user=master.user_for_token(master.login("admin","secret")["token"])
            manager=CompanyManager(Path(folder)/"master.db")
            company=manager.create_company({"name":"Second Company","year":2026,"mof_number":"MOF-2"},master)
            db_2026=manager.database(company["id"],2026)
            db_2026.save_party({"name":"Only In Company Two","kind":"supplier","currency":"USD"},user["id"])
            self.assertFalse(any(p["name"]=="Only In Company Two" for p in master.list_parties()))
            manager.create_year(company["id"],2027,user["id"])
            db_2027=manager.database(company["id"],2027)
            self.assertNotEqual(Path(db_2026.path).resolve(),Path(db_2027.path).resolve())
            self.assertTrue(any(p["name"]=="Only In Company Two" for p in db_2027.list_parties()))
            updated=manager.update_company(company["id"],{"name":"Renamed Company","active":False})
            self.assertEqual(updated["name"],"Renamed Company"); self.assertFalse(updated["active"])

    def test_company_closing_creates_new_year_and_opening_journal_voucher(self):
        with tempfile.TemporaryDirectory() as folder:
            master=Database(Path(folder)/"master.db"); master.initialize("secret")
            user=master.user_for_token(master.login("admin","secret")["token"])
            manager=CompanyManager(Path(folder)/"master.db")
            company=manager.create_company({"name":"Closing Test","year":2026},master)
            source=manager.database(company["id"],2026)
            source.import_invoice({"invoice_number":"S-1","invoice_date":"30-12-2026","party_name":"Client",
                "kind":"sale","currency":"USD","subtotal":1000,"vat":110,"total":1110},user["id"])
            result=manager.close_and_open_year(company["id"],2026,user["id"])
            self.assertEqual(result["opened_year"],2027)
            self.assertIn("OPEN-2027-USD",result["opening_vouchers"])
            target=manager.database(company["id"],2027)
            opening=[row for row in target.journal() if row["source_type"]=="opening"]
            self.assertTrue(opening)
            self.assertAlmostEqual(sum(row["debit"] for row in opening),sum(row["credit"] for row in opening),places=2)
            self.assertEqual(target.profit_and_loss("2027-01-01","2027-12-31","USD"),[])
            years={int(item["year"]):item["status"] for item in manager._company(company["id"])["years"]}
            self.assertEqual(years[2026],"closed"); self.assertEqual(years[2027],"open")

    def test_separate_year_does_not_close_previous_and_admin_can_reopen(self):
        with tempfile.TemporaryDirectory() as folder:
            master=Database(Path(folder)/"master.db"); master.initialize("secret")
            user=master.user_for_token(master.login("admin","secret")["token"])
            manager=CompanyManager(Path(folder)/"master.db")
            company=manager.create_company({"name":"Fiscal Control","year":2024},master)
            manager.create_year(company["id"],2025,user["id"])
            self.assertEqual(manager.year_status(company["id"],2024),"open")
            self.assertEqual(manager.year_status(company["id"],2025),"open")
            source=manager.database(company["id"],2024)
            source.import_invoice({"invoice_number":"Y24-1","invoice_date":"30-12-2024","party_name":"Client 2024","kind":"sale","currency":"USD","subtotal":100,"vat":11,"total":111},user["id"])
            refreshed=manager.refresh_opening(company["id"],2024,user["id"])
            self.assertTrue(refreshed["provisional"]); self.assertIn("OPEN-2025-USD",refreshed["opening_vouchers"])
            self.assertTrue(any(row["source_type"]=="invoice" for row in source.journal()))
            closed=manager.close_and_open_year(company["id"],2024,user["id"])
            self.assertEqual(closed["closed_year"],2024)
            self.assertEqual(manager.year_status(company["id"],2024),"closed")
            reopened=manager.reopen_year(company["id"],2024,user["id"])
            self.assertEqual(reopened["status"],"open")
            self.assertEqual(manager.year_status(company["id"],2024),"open")

    def test_opening_display_currency_dc_choices_and_linked_rate(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"currency_views.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            db.save_exchange_rate({"date_from":"01-01-2026","date_to":"31-12-2026","from_currency":"USD","to_currency":"LBP","rate":"89500"},user["id"])
            db.save_exchange_rate({"date_from":"01-01-2026","date_to":"31-12-2026","from_currency":"EUR","to_currency":"USD","rate":"1.10"},user["id"])
            with db.connect() as connection:
                linked=connection.execute("SELECT CAST(rate AS REAL) rate FROM exchange_rates WHERE rate_date='20-09-2026' AND from_currency='EUR' AND to_currency='LBP'").fetchone()
            self.assertAlmostEqual(linked["rate"],98450)
            invoice_id=db.import_invoice({"invoice_number":"OPEN-1","invoice_date":"01-09-2026","party_name":"Supplier X","kind":"purchases","currency":"USD","deductible_subtotal":100,"non_deductible_subtotal":20,"vat":11,"total":131,"expense_no_vat_side":"C - Credit"},user["id"])
            party=next(p for p in db.list_parties() if p["name"]=="Supplier X")
            with_opening=db.statement_of_account(party["id"],"2026-09-10","2026-09-30",None,True,"LBP")
            without_opening=db.statement_of_account(party["id"],"2026-09-10","2026-09-30",None,False,"LBP")
            self.assertTrue(with_opening["opening"]); self.assertFalse(without_opening["opening"])
            expense_id=db.add_expense({"expense_date":"20-09-2026","description":"Credit-side correction","currency":"USD","with_vat_subtotal":10,"without_vat_subtotal":5,"vat":1.1,"expense_account":"601100000","expense_without_vat_account":"601100001","vat_account":"442660000","payment_account":"531","expense_side":"C - Credit","expense_without_vat_side":"D - Debit","vat_side":"D - Debit","payment_side":"D - Debit"},user["id"])
            lines=[r for r in db.journal(currency="USD") if r["source_type"]=="expense" and r["source_id"]==expense_id]
            self.assertEqual(next(r for r in lines if r["account_code"]=="601100000")["credit"],10)
            trial=db.trial_balance("2026-09-20","2026-09-20")
            self.assertTrue(all("usd_balance" in r and "lbp_balance" in r for r in trial))

    def test_payroll_employee_number_calculation_and_editable_rules(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/"payroll.db"); db.initialize("secret")
            user=db.user_for_token(db.login("admin","secret")["token"])
            employee=db.save_employee({"employee_number":"1234","full_name":"Payroll Employee","currency":"LBP",
                "base_salary":"100000000","marital_status":"married","children":1},user["id"])
            self.assertEqual(employee["employee_number"],"123400001")
            second=db.save_employee({"employee_number":"1234","full_name":"Second Employee","currency":"LBP","base_salary":"1"},user["id"])
            self.assertEqual(second["employee_number"],"123400002")
            settings=db.payroll_settings_for("2025-06-30")
            settings.update({"date_from":"2025-01-01","employee_ceiling":"50000000","medical_ceiling":"60000000",
                "family_ceiling":"70000000","end_service_ceiling":"0"})
            db.save_payroll_settings(settings,user["id"])
            result=db.calculate_payroll({"employee_id":employee["id"],"period_date":"2025-06-30","overtime":"5000000",
                "transport":"10000000","commission":"2000000","schooling":"3000000","bonus":"4000000","thirteenth_month":"0"})
            self.assertEqual(result["gross_salary"],124000000)
            self.assertEqual(result["employee_nssf"],1500000)
            self.assertGreater(result["income_tax"],0)
            employee["spouse_works"]=True; db.save_employee(employee,user["id"])
            working_spouse=db.calculate_payroll({"employee_id":employee["id"],"period_date":"2025-06-30","overtime":"5000000",
                "transport":"10000000","commission":"2000000","schooling":"3000000","bonus":"4000000","thirteenth_month":"0"})
            self.assertGreater(working_spouse["income_tax"],result["income_tax"])
            self.assertEqual(working_spouse["income_tax_lbp"],working_spouse["income_tax"])
            usd_employee=db.save_employee({"employee_number":"5555","full_name":"USD Employee","currency":"USD","base_salary":"2000","marital_status":"single"},user["id"])
            usd_result=db.calculate_payroll({"employee_id":usd_employee["id"],"period_date":"2025-06-30"})
            self.assertGreater(usd_result["income_tax_lbp"],0)
            self.assertAlmostEqual(usd_result["income_tax"]*89500,usd_result["income_tax_lbp"],delta=1000)
            self.assertEqual(usd_result["employee_nssf_lbp"],1500000)
            retro=db.calculate_payroll({"employee_id":employee["id"],"period_date":"2025-06-30","retro_salary":"10000000"})
            self.assertEqual(retro["gross_salary"],110000000)
            self.assertGreater(retro["income_tax"],0)
            saved=db.save_payroll({"employee_id":employee["id"],"period_date":"2025-06-30","overtime":"5000000",
                "transport":"10000000","commission":"2000000","schooling":"3000000","bonus":"4000000"},user["id"])
            self.assertTrue(saved["payroll_number"].startswith("PAY-202506-"))
            self.assertEqual(saved["status"],"draft")
            posted=db.post_payroll(saved["id"],user["id"])
            self.assertEqual(posted["status"],"posted")
            lines=[row for row in db.journal() if row["source_type"]=="payroll" and row["source_id"]==saved["id"]]
            self.assertTrue(lines)
            self.assertAlmostEqual(sum(row["debit"] for row in lines),sum(row["credit"] for row in lines),places=2)
            with self.assertRaisesRegex(ValueError,"already posted"): db.post_payroll(saved["id"],user["id"])

if __name__ == "__main__": unittest.main()
