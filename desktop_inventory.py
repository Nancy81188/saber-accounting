"""Inventory screens (version 1.20)."""
from __future__ import annotations

import tkinter as tk
from datetime import datetime
from tkinter import filedialog, messagebox, ttk

from desktop_brains import EditableSheet

NAVY, GOLD, LIGHT = "#071b2e", "#c9a96a", "#f3f6f8"
RED, MUTED = "#8B1E1E", "#5f6b76"
DOC_TYPES = {"Opening Stock": "opening", "Stock Receipt": "receipt", "Stock Issue": "issue", "Adjustment +": "adjustment_in", "Adjustment -": "adjustment_out", "Transfer": "transfer"}
REPORTS = {"Stock Valuation": "valuation", "Stock Card": "stock_card", "Stock Movements": "movements", "Sales Margin (COGS)": "margin", "Reorder Report": "reorder", "Slow-moving Stock": "slow"}


def _num(value):
    try: return float(str(value or 0).replace(",", ""))
    except ValueError: return None


def _dd(value):
    text = str(value or "")
    return f"{text[8:10]}-{text[5:7]}-{text[:4]}" if len(text) == 10 and text[4] == "-" else text


class InventoryMixin:
    def build_inventory(self):
        nested = ttk.Notebook(self.inventory_tab); nested.pack(fill="both", expand=True, padx=8, pady=8)
        pages = {name: tk.Frame(nested, bg=LIGHT) for name in ("Items", "Categories & Units", "Stock In / Stock Out", "Stock Documents", "Physical Inventory", "Inventory Reports", "Warehouses & Settings")}
        for name, page in pages.items(): nested.add(page, text=name)
        self.build_items_page(pages["Items"]); self.build_categories_page(pages["Categories & Units"]); self.build_stock_in_out_page(pages["Stock In / Stock Out"])
        self.build_physical_page(pages["Physical Inventory"]); self.build_stock_documents_page(pages["Stock Documents"])
        self.build_inventory_reports_page(pages["Inventory Reports"]); self.build_inventory_settings_page(pages["Warehouses & Settings"])
        self.load_inventory()

    def load_inventory(self):
        if not hasattr(self, "items_tree") or not self.items_tree.winfo_exists(): return
        try: self.inventory_rows = self.client.inventory_items(); self.warehouse_rows = self.client.warehouses(); inv = self.client.inventory_settings()
        except Exception as exc: return messagebox.showerror("Inventory", str(exc))
        self.inventory_currency = inv["currency"]
        self.items_tree.delete(*self.items_tree.get_children())
        for i in self.inventory_rows:
            status = "Inactive" if not i["active"] else "Reorder" if i["reorder_level"] and i["quantity"] <= i["reorder_level"] else "OK"
            self.items_tree.insert("", "end", iid=str(i["id"]), values=(i["sku"], i["name"], i.get("category") or "", i.get("subcategory") or "", i.get("supplier_name") or "", i["unit"],
                f'{i["quantity"]:,.3f}', f'{i["average_cost"]:,.4f}', f'{i["stock_value"]:,.2f}', f'{i["sales_price"]:,.2f}', status), tags=("reorder",) if status == "Reorder" else ())
        self.load_item_lists()
        names = [f'{w["code"]} - {w["name"]}' for w in self.warehouse_rows if w["active"]]
        for box, var in ((getattr(self, "sio_wh_box", None), getattr(self, "sio_vars", {}).get("warehouse")), (getattr(self, "pc_wh_box", None), getattr(self, "pc_vars", {}).get("warehouse"))):
            if box is not None and box.winfo_exists():
                box["values"] = names
                if var is not None and not var.get() and names: var.set(names[0])
        if hasattr(self, "pc_find_box"): self.load_counts()
        for box in (getattr(self, "sd_warehouse_box", None), getattr(self, "sd_to_box", None)):
            if box is not None and box.winfo_exists(): box["values"] = names
        if hasattr(self, "ir_warehouse_box"): self.ir_warehouse_box["values"] = ["All"] + names
        items = [f'{i["sku"]} - {i["name"]}' for i in self.inventory_rows]
        if hasattr(self, "ir_item_box"): self.ir_item_box["values"] = [""] + items
        categories = sorted({i.get("category") for i in self.inventory_rows if i.get("category")})
        if hasattr(self, "ir_category_box"): self.ir_category_box["values"] = ["All"] + categories
        if hasattr(self, "warehouses_tree"):
            self.warehouses_tree.delete(*self.warehouses_tree.get_children())
            for w in self.warehouse_rows: self.warehouses_tree.insert("", "end", iid=str(w["id"]), values=(w["code"], w["name"], "Yes" if w["active"] else "No"))
            self.inv_currency.set(inv["currency"]); self.inv_method.set("FIFO" if inv["method"] == "fifo" else "Weighted average")
        self.load_stock_documents()

    def item_by_code(self, code):
        code = str(code or "").split(" - ", 1)[0].strip().upper()
        if not code: return None
        return next((i for i in getattr(self, "inventory_rows", []) if i["sku"].upper() == code or (i.get("barcode") or "").upper() == code), None)

    # ------------------------------------------------------------ items
    def build_items_page(self, page):
        form = tk.LabelFrame(page, text="Item", bg=LIGHT, padx=8, pady=5); form.pack(fill="x", padx=8, pady=6)
        self.item_id = None; self.item_vars = {k: tk.StringVar() for k in ("sku", "name", "unit", "category", "subcategory", "supplier_name", "location", "sales_price", "reorder_level", "barcode", "notes", "default_vat", "cost_account")}
        self.item_vars["unit"].set("unit"); self.item_vars["default_vat"].set("11%"); self.item_active = tk.BooleanVar(value=True)
        self.item_boxes = {}
        for index, (key, label, width) in enumerate((("sku", "Item Code (auto if blank)", 14), ("name", "Item Name", 24), ("unit", "Unit", 12), ("category", "Category", 16),
                                                     ("subcategory", "Subcategory", 16), ("supplier_name", "Supplier", 22), ("sales_price", "Sales Price", 11), ("default_vat", "Default VAT", 9), ("reorder_level", "Reorder Level", 9),
                                                     ("location", "Location (shelf)", 12), ("barcode", "Barcode", 14), ("notes", "Notes", 24), ("cost_account", "Cost Account (opt.)", 14))):
            tk.Label(form, text=label, bg=LIGHT).grid(row=index // 3, column=(index % 3) * 2, sticky="w", padx=4, pady=2)
            if key == "default_vat":
                widget = ttk.Combobox(form, textvariable=self.item_vars[key], values=["11%", "0%"], state="readonly", width=width); self.item_boxes[key] = widget
            elif key in ("unit", "category", "subcategory", "supplier_name"):
                widget = ttk.Combobox(form, textvariable=self.item_vars[key], width=width); self.item_boxes[key] = widget
                if key == "category": widget.bind("<<ComboboxSelected>>", lambda _e: self.item_category_chosen())
            else: widget = tk.Entry(form, textvariable=self.item_vars[key], width=width)
            widget.grid(row=index // 3, column=(index % 3) * 2 + 1, sticky="w", padx=4, pady=2)
        self.item_cost_label = tk.Label(form, text="Cost price (average of purchases): -", bg=LIGHT, fg=NAVY, font=("Segoe UI", 9, "bold")); self.item_cost_label.grid(row=3, column=4, columnspan=2, sticky="w", padx=4)
        buttons = tk.Frame(form, bg=LIGHT); buttons.grid(row=5, column=0, columnspan=6, sticky="w", pady=(4, 0))
        tk.Checkbutton(buttons, text="Active", variable=self.item_active, bg=LIGHT).pack(side="left", padx=(0, 8))
        self.action_button(buttons, "New", self.new_item).pack(side="left", padx=3)
        tk.Button(buttons, text="Save", command=self.save_item, bg=GOLD, fg=NAVY, border=0, padx=18, pady=6, font=("Segoe UI", 9, "bold")).pack(side="left", padx=3)
        self.action_button(buttons, "Stock Card", lambda: self.open_stock_card()).pack(side="left", padx=3)
        tk.Label(buttons, text="Quantity and cost come from the stock documents. Double-click an item to edit it.", bg=LIGHT, fg=MUTED).pack(side="left", padx=10)
        self.items_tree = self.table(page, [("sku", "Item Code", 95), ("name", "Item", 200), ("category", "Category", 95), ("subcategory", "Subcategory", 95), ("supplier", "Supplier", 120),
            ("unit", "Unit", 50), ("qty", "On Hand", 80), ("cost", "Cost Price (avg)", 100), ("value", "Stock Value", 100), ("price", "Sales Price", 85), ("status", "Status", 65)])
        self.items_tree.tag_configure("reorder", foreground=RED); self.items_tree.bind("<Double-1>", lambda _e: self.edit_item())

    def new_item(self):
        self.item_id = None; [v.set("") for v in self.item_vars.values()]; self.item_vars["unit"].set("unit"); self.item_vars["default_vat"].set("11%"); self.item_active.set(True)
        if hasattr(self, "item_cost_label"): self.item_cost_label.config(text="Cost price (average of purchases): -")

    def edit_item(self):
        selected = self.items_tree.selection()
        if not selected: return
        item = next(i for i in self.inventory_rows if str(i["id"]) == selected[0]); self.item_id = item["id"]
        for key in self.item_vars: self.item_vars[key].set("" if item.get(key) in (None, 0.0) and key in ("barcode", "notes", "category") else str(item.get(key) if item.get(key) is not None else ""))
        self.item_vars["default_vat"].set("0%" if str(item.get("default_vat") or "11").strip() in ("0", "0.0", "0%") else "11%")
        self.item_vars["sales_price"].set(f'{item["sales_price"]:g}'); self.item_vars["reorder_level"].set(f'{item["reorder_level"]:g}'); self.item_active.set(bool(item["active"]))
        self.item_cost_label.config(text=f'Cost price (average of purchases): {item["average_cost"]:,.4f} {getattr(self, "inventory_currency", "")}   On hand: {item["quantity"]:,.3f}')

    def save_item(self):
        payload = {k: v.get().strip() for k, v in self.item_vars.items()}; payload.update(id=self.item_id, active=self.item_active.get())
        payload["default_vat"] = "0" if payload.get("default_vat", "11%").replace("%", "").strip() in ("0", "0.0") else "11"
        try: saved = self.client.save_inventory_item(payload)
        except Exception as exc: return messagebox.showerror("Items", str(exc))
        self.new_item(); self.load_inventory(); messagebox.showinfo("Items", f'Item {saved["sku"]} - {saved["name"]} saved')

    def open_stock_card(self):
        selected = self.items_tree.selection()
        if not selected: return messagebox.showwarning("Stock Card", "Select an item first")
        item = next(i for i in self.inventory_rows if str(i["id"]) == selected[0])
        self.ir_report.set("Stock Card"); self.ir_item.set(f'{item["sku"]} - {item["name"]}')
        self.inventory_notebook_select("Inventory Reports"); self.run_inventory_report()

    def inventory_notebook_select(self, name):
        notebook = [w for w in self.inventory_tab.winfo_children() if isinstance(w, ttk.Notebook)][0]
        notebook.select([t for t in notebook.tabs() if notebook.tab(t, "text") == name][0])

    # ------------------------------------------------------------ stock documents
    def build_stock_documents_page(self, page):
        self.sd_id = None; self.sd_vars = {k: tk.StringVar() for k in ("type", "number", "date", "warehouse", "to_warehouse", "party", "reference", "notes", "find")}
        v = self.sd_vars; v["type"].set("Stock Receipt"); v["date"].set(datetime.now().strftime("%d-%m-%Y"))
        bar = tk.Frame(page, bg=LIGHT); bar.pack(fill="x", padx=8, pady=(6, 2))
        tk.Label(bar, text="Type", bg=LIGHT, font=("Segoe UI", 9, "bold")).pack(side="left")
        type_box = ttk.Combobox(bar, textvariable=v["type"], values=list(DOC_TYPES), state="readonly", width=14); type_box.pack(side="left", padx=(4, 8))
        type_box.bind("<<ComboboxSelected>>", lambda _e: self.stock_type_changed())
        tk.Label(bar, text="Number", bg=LIGHT).pack(side="left")
        tk.Entry(bar, textvariable=v["number"], width=16, state="readonly", readonlybackground="white", font=("Segoe UI", 10, "bold")).pack(side="left", padx=(4, 8))
        tk.Label(bar, text="Date", bg=LIGHT).pack(side="left"); self.date_entry(bar, v["date"], 11).pack(side="left", padx=(4, 8))
        tk.Label(bar, text="Warehouse", bg=LIGHT).pack(side="left")
        self.sd_warehouse_box = ttk.Combobox(bar, textvariable=v["warehouse"], state="readonly", width=17); self.sd_warehouse_box.pack(side="left", padx=(4, 8))
        self.sd_to_label = tk.Label(bar, text="To", bg=LIGHT); self.sd_to_box = ttk.Combobox(bar, textvariable=v["to_warehouse"], state="readonly", width=17)
        tk.Label(bar, text="Find", bg=LIGHT).pack(side="right")
        self.sd_find_box = ttk.Combobox(bar, textvariable=v["find"], width=26); self.sd_find_box.pack(side="right", padx=4)
        self.sd_find_box.bind("<<ComboboxSelected>>", lambda _e: self.open_stock_document())
        bar2 = tk.Frame(page, bg=LIGHT); bar2.pack(fill="x", padx=8, pady=2)
        tk.Label(bar2, text="Customer / Supplier", bg=LIGHT).pack(side="left")
        self.sd_party_box = ttk.Combobox(bar2, textvariable=v["party"], width=26); self.sd_party_box.pack(side="left", padx=(4, 8))
        tk.Label(bar2, text="Reference (invoice / PO)", bg=LIGHT).pack(side="left"); tk.Entry(bar2, textvariable=v["reference"], width=16).pack(side="left", padx=(4, 8))
        tk.Label(bar2, text="Notes", bg=LIGHT).pack(side="left"); tk.Entry(bar2, textvariable=v["notes"], width=28).pack(side="left", padx=4)
        bottom = tk.Frame(page, bg=LIGHT); bottom.pack(side="bottom", fill="x", padx=8, pady=6)
        self.action_button(bottom, "New", self.new_stock_document).pack(side="left", padx=(0, 3))
        self.action_button(bottom, "Add Line", self.add_stock_line).pack(side="left", padx=3)
        tk.Button(bottom, text="Delete Line", command=self.delete_stock_line, bg=RED, fg="white", border=0, padx=12, pady=7).pack(side="left", padx=3)
        tk.Button(bottom, text="Save", command=self.save_stock_document, bg=GOLD, fg=NAVY, border=0, padx=18, pady=7, font=("Segoe UI", 9, "bold")).pack(side="left", padx=(12, 3))
        tk.Button(bottom, text="Delete Document", command=self.delete_stock_document, bg=RED, fg="white", border=0, padx=12, pady=7).pack(side="left", padx=3)
        self.sd_total = tk.Label(bottom, text="", bg=LIGHT, fg=NAVY, font=("Segoe UI", 10, "bold")); self.sd_total.pack(side="right", padx=8)
        self.sd_info = tk.Label(page, text="Item Code: type the code or press F2 in the cell. Issues, adjustments - and transfers leave at the current cost; receipts and opening stock need a unit cost.",
                                bg=LIGHT, fg=MUTED, anchor="w"); self.sd_info.pack(side="bottom", fill="x", padx=10)
        columns = [("line", "#", 40, "center"), ("sku", "Item Code", 110, "w"), ("name", "Item", 260, "w"), ("unit", "Unit", 60, "center"), ("quantity", "Quantity", 100, "e"),
                   ("unit_cost", "Unit Cost", 110, "e"), ("value", "Value", 120, "e"), ("on_hand", "On Hand", 100, "e")]
        self.stock_sheet = EditableSheet(self, page, columns, ["sku", "quantity", "unit_cost"], self.stock_cell_changed, height=9)
        self.stock_sheet.tree.bind("<F2>", lambda _e: self.stock_item_lookup())
        self.new_stock_document()

    def stock_item_lookup(self):
        iid, row = self.stock_sheet.selected()
        if not row: return
        window = tk.Toplevel(self); window.title("Items - F2"); window.geometry("640x420"); window.configure(bg=LIGHT); window.transient(self); window.grab_set()
        search = tk.StringVar(); entry = tk.Entry(window, textvariable=search, width=40); entry.pack(padx=10, pady=8); entry.focus_set()
        tree = ttk.Treeview(window, columns=("sku", "name", "qty"), show="headings"); tree.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        for key, label, width in (("sku", "Code", 110), ("name", "Item", 320), ("qty", "On Hand", 100)): tree.heading(key, text=label); tree.column(key, width=width)
        def fill(*_a):
            tree.delete(*tree.get_children()); text = search.get().casefold()
            for i in self.inventory_rows:
                if i["active"] and (not text or text in f'{i["sku"]} {i["name"]}'.casefold()): tree.insert("", "end", values=(i["sku"], i["name"], f'{i["quantity"]:,.3f}'))
        def choose(_e=None):
            if tree.selection(): self.stock_cell_changed(iid, "sku", tree.item(tree.selection()[0], "values")[0]); self.stock_sheet.refresh(iid); window.destroy()
        search.trace_add("write", fill); tree.bind("<Double-1>", choose); tree.bind("<Return>", choose); fill()

    def stock_type_changed(self):
        transfer = self.sd_vars["type"].get() == "Transfer"
        if transfer: self.sd_to_label.pack(side="left"); self.sd_to_box.pack(side="left", padx=(4, 8))
        else: self.sd_to_label.pack_forget(); self.sd_to_box.pack_forget()
        if not self.sd_id:
            try: self.sd_vars["number"].set(self.client.next_stock_number(DOC_TYPES[self.sd_vars["type"].get()], self.sd_vars["date"].get()))
            except Exception: self.sd_vars["number"].set("")
        for iid, row in self.stock_sheet.rows.items(): self.refresh_stock_row(row); self.stock_sheet.refresh(iid)
        self.update_stock_total()

    def cost_is_automatic(self):
        return DOC_TYPES[self.sd_vars["type"].get()] in ("issue", "adjustment_out", "transfer")

    def refresh_stock_row(self, row):
        item = self.item_by_code(row.get("sku"))
        if item:
            row["name"] = item["name"]; row["unit"] = item["unit"]; row["on_hand"] = f'{item["quantity"]:,.3f}'
            if self.cost_is_automatic() or not row.get("unit_cost"): row["unit_cost"] = round(item["average_cost"], 4) if self.cost_is_automatic() or item["average_cost"] else row.get("unit_cost", "")
        qty = _num(row.get("quantity")) or 0; cost = _num(row.get("unit_cost")) or 0
        row["value"] = f"{qty * cost:,.2f}" if qty else ""
        row["_display"] = {"quantity": f"{qty:,.3f}" if qty else "", "unit_cost": f"{cost:,.4f}" if row.get("unit_cost") not in ("", None) else ""}

    def stock_cell_changed(self, iid, key, text):
        row = self.stock_sheet.rows[iid]
        if key == "sku":
            item = self.item_by_code(text)
            if text and not item: messagebox.showwarning("Stock Documents", f"Item {text} was not found. Press F2 on the line to search."); return False
            row["sku"] = item["sku"] if item else ""
        else:
            value = _num(text)
            if value is None or value < 0: messagebox.showwarning("Stock Documents", "Enter a positive number"); return False
            if key == "unit_cost" and self.cost_is_automatic(): messagebox.showinfo("Stock Documents", "Issues, adjustments - and transfers use the current average cost automatically"); return False
            row[key] = value
        self.refresh_stock_row(row); self.update_stock_total()
        rows = self.stock_sheet.tree.get_children()
        if rows and iid == rows[-1] and row.get("sku") and row.get("quantity"): self.add_stock_line(edit=False)

    def add_stock_line(self, edit=True):
        row = {"sku": "", "quantity": "", "unit_cost": ""}; self.refresh_stock_row(row); iid = self.stock_sheet.insert(row)
        self.stock_sheet.tree.selection_set(iid); self.stock_sheet.tree.focus(iid)
        if edit: self.after(30, lambda: self.stock_sheet.edit(iid, "sku"))

    def delete_stock_line(self):
        if not self.stock_sheet.delete_selected(): return messagebox.showwarning("Stock Documents", "Select a line first")
        self.update_stock_total()

    def update_stock_total(self):
        rows = [r for r in self.stock_sheet.ordered() if r.get("sku") and _num(r.get("quantity"))]
        total = sum((_num(r["quantity"]) or 0) * (_num(r.get("unit_cost")) or 0) for r in rows)
        self.sd_total.config(text=f"{len(rows)} line(s)   Total value: {total:,.2f} {getattr(self, 'inventory_currency', '')}")

    def new_stock_document(self):
        self.sd_id = None; v = self.sd_vars
        for key in ("party", "reference", "notes", "find"): v[key].set("")
        v["date"].set(datetime.now().strftime("%d-%m-%Y"))
        if not v["warehouse"].get() and getattr(self, "warehouse_rows", None): v["warehouse"].set(f'{self.warehouse_rows[0]["code"]} - {self.warehouse_rows[0]["name"]}')
        self.stock_sheet.clear(); self.add_stock_line(edit=False); self.add_stock_line(edit=False); self.stock_type_changed()

    def load_stock_documents(self):
        if not hasattr(self, "sd_find_box"): return
        try: docs = self.client.stock_documents(); parties = self.client.parties()
        except Exception: return
        labels = {v: k for k, v in DOC_TYPES.items()}
        self.sd_doc_map = {f'{d["number"]} | {_dd(d["doc_date"])} | {labels[d["doc_type"]]} | {d.get("party_name") or d.get("reference") or ""}': d["id"] for d in docs}
        self.sd_find_box["values"] = list(self.sd_doc_map)
        self.sd_party_map = {p["name"]: p for p in parties}; self.sd_party_box["values"] = list(self.sd_party_map)

    def open_stock_document(self):
        document_id = getattr(self, "sd_doc_map", {}).get(self.sd_vars["find"].get())
        if not document_id: return
        try: doc = self.client.stock_document(document_id)
        except Exception as exc: return messagebox.showerror("Stock Documents", str(exc))
        self.sd_id = doc["id"]; v = self.sd_vars; labels = {val: k for k, val in DOC_TYPES.items()}
        v["type"].set(labels[doc["doc_type"]]); v["number"].set(doc["number"]); v["date"].set(_dd(doc["doc_date"]))
        v["warehouse"].set(f'{doc["warehouse_code"]} - {doc["warehouse_name"]}'); v["party"].set(doc.get("party_name") or ""); v["reference"].set(doc.get("reference") or ""); v["notes"].set(doc.get("notes") or "")
        if doc.get("to_warehouse_code"): v["to_warehouse"].set(next((f'{w["code"]} - {w["name"]}' for w in self.warehouse_rows if w["code"] == doc["to_warehouse_code"]), ""))
        self.stock_sheet.clear()
        for line in doc["lines"]:
            row = {"sku": line["sku"], "quantity": line["quantity"], "unit_cost": line["unit_cost"]}; self.refresh_stock_row(row); self.stock_sheet.insert(row)
        self.stock_type_changed(); v["number"].set(doc["number"])
        self.sd_info.config(text=f"Document {doc['number']} opened" + (" (created by an invoice - change the invoice instead)" if doc.get("invoice_id") else ""))

    def save_stock_document(self):
        v = self.sd_vars; lines = [r for r in self.stock_sheet.ordered() if r.get("sku") and _num(r.get("quantity"))]
        if not lines: return messagebox.showwarning("Stock Documents", "Add at least one item with a quantity")
        doc_type = DOC_TYPES[v["type"].get()]
        if doc_type in ("opening", "receipt") and any(not _num(r.get("unit_cost")) for r in lines): return messagebox.showwarning("Stock Documents", "Enter the unit cost of every line")
        party = getattr(self, "sd_party_map", {}).get(v["party"].get())
        header = {"doc_type": doc_type, "doc_date": v["date"].get().strip(), "warehouse_id": v["warehouse"].get().split(" - ", 1)[0], "to_warehouse_id": v["to_warehouse"].get().split(" - ", 1)[0],
                  "party_id": party["id"] if party else None, "reference": v["reference"].get().strip(), "notes": v["notes"].get().strip()}
        payload = [{"sku": r["sku"], "quantity": r["quantity"], "unit_cost": 0 if self.cost_is_automatic() else r.get("unit_cost")} for r in lines]
        try: saved = self.client.save_stock_document(header, payload, self.sd_id)
        except Exception as exc: return messagebox.showerror("Stock Documents", str(exc))
        messagebox.showinfo("Stock Documents", f'{saved["number"]} saved'); self.load_inventory(); self.new_stock_document()

    def delete_stock_document(self):
        if not self.sd_id: return messagebox.showwarning("Stock Documents", "Open a saved document first (Find)")
        if not messagebox.askyesno("Stock Documents", f"Delete {self.sd_vars['number'].get()}?"): return
        try: self.client.delete_stock_document(self.sd_id)
        except Exception as exc: return messagebox.showerror("Stock Documents", str(exc))
        self.load_inventory(); self.new_stock_document()

    # ------------------------------------------------------------ reports
    def build_inventory_reports_page(self, page):
        bar = tk.Frame(page, bg=LIGHT); bar.pack(fill="x", padx=8, pady=6); year = getattr(self, "current_fiscal_year", datetime.now().year)
        self.ir_report = tk.StringVar(value="Stock Valuation"); self.ir_from = tk.StringVar(value=f"01-01-{year}"); self.ir_to = tk.StringVar(value=f"31-12-{year}")
        self.ir_warehouse = tk.StringVar(value="All"); self.ir_item = tk.StringVar(); self.ir_method = tk.StringVar(value="Company setting"); self.ir_days = tk.StringVar(value="90")
        self.ir_category = tk.StringVar(value="All"); self.ir_zero = tk.BooleanVar(value=False)
        ttk.Combobox(bar, textvariable=self.ir_report, values=list(REPORTS), state="readonly", width=20).pack(side="left", padx=(0, 8))
        tk.Label(bar, text="From", bg=LIGHT).pack(side="left"); self.date_entry(bar, self.ir_from, 11).pack(side="left", padx=(4, 6))
        tk.Label(bar, text="To / As of", bg=LIGHT).pack(side="left"); self.date_entry(bar, self.ir_to, 11).pack(side="left", padx=(4, 6))
        tk.Label(bar, text="Warehouse", bg=LIGHT).pack(side="left")
        self.ir_warehouse_box = ttk.Combobox(bar, textvariable=self.ir_warehouse, state="readonly", width=16); self.ir_warehouse_box.pack(side="left", padx=(4, 6))
        tk.Label(bar, text="Costing", bg=LIGHT).pack(side="left")
        ttk.Combobox(bar, textvariable=self.ir_method, values=["Company setting", "Weighted average", "FIFO"], state="readonly", width=15).pack(side="left", padx=4)
        bar2 = tk.Frame(page, bg=LIGHT); bar2.pack(fill="x", padx=8)
        tk.Label(bar2, text="Item (Stock Card)", bg=LIGHT).pack(side="left")
        self.ir_item_box = ttk.Combobox(bar2, textvariable=self.ir_item, width=30); self.ir_item_box.pack(side="left", padx=(4, 8))
        tk.Label(bar2, text="Category", bg=LIGHT).pack(side="left")
        self.ir_category_box = ttk.Combobox(bar2, textvariable=self.ir_category, state="readonly", width=14); self.ir_category_box.pack(side="left", padx=(4, 8))
        self.ir_subcategory = tk.StringVar(value="All"); self.ir_unit = tk.StringVar(value="All"); self.ir_supplier = tk.StringVar(value="All")
        bar3 = tk.Frame(page, bg=LIGHT); bar3.pack(fill="x", padx=8, pady=(4, 0), after=bar2)
        tk.Label(bar3, text="Subcategory", bg=LIGHT).pack(side="left"); self.ir_subcategory_box = ttk.Combobox(bar3, textvariable=self.ir_subcategory, state="readonly", width=14); self.ir_subcategory_box.pack(side="left", padx=(4, 8))
        tk.Label(bar3, text="Unit", bg=LIGHT).pack(side="left"); self.ir_unit_box = ttk.Combobox(bar3, textvariable=self.ir_unit, state="readonly", width=8); self.ir_unit_box.pack(side="left", padx=(4, 8))
        tk.Label(bar3, text="Supplier", bg=LIGHT).pack(side="left"); self.ir_supplier_box = ttk.Combobox(bar3, textvariable=self.ir_supplier, state="readonly", width=22); self.ir_supplier_box.pack(side="left", padx=(4, 8))
        tk.Label(bar3, text="Filters combine: e.g. Category + Supplier + Warehouse", bg=LIGHT, fg=MUTED).pack(side="left", padx=6)
        tk.Label(bar2, text="Slow-moving days", bg=LIGHT).pack(side="left"); tk.Entry(bar2, textvariable=self.ir_days, width=5).pack(side="left", padx=4)
        tk.Checkbutton(bar2, text="Include zero stock", variable=self.ir_zero, bg=LIGHT).pack(side="left", padx=6)
        tk.Button(bar2, text="Show", command=self.run_inventory_report, bg=GOLD, fg=NAVY, border=0, padx=18, pady=6, font=("Segoe UI", 9, "bold")).pack(side="left", padx=6)
        for text, fmt in (("Print", "print"), ("Excel", "xlsx"), ("PDF", "pdf")): self.action_button(bar2, text, lambda f=fmt: self.export_inventory_report(f)).pack(side="left", padx=2)
        self.ir_info = tk.Label(page, text="", bg=LIGHT, fg=NAVY, anchor="w"); self.ir_info.pack(fill="x", padx=10)
        self.ir_viewer = self.report_viewer(page)

    def inventory_report_options(self):
        options = {"date_from": self.ir_from.get().strip(), "date_to": self.ir_to.get().strip(), "days": self.ir_days.get().strip() or "90", "include_zero": self.ir_zero.get()}
        if self.ir_warehouse.get() not in ("", "All"):
            code = self.ir_warehouse.get().split(" - ", 1)[0]; options["warehouse_id"] = next(w["id"] for w in self.warehouse_rows if w["code"] == code)
        item = self.item_by_code(self.ir_item.get())
        if item: options["item_id"] = item["id"]
        if self.ir_category.get() not in ("", "All"): options["category"] = self.ir_category.get()
        if self.ir_subcategory.get() not in ("", "All"): options["subcategory"] = self.ir_subcategory.get()
        if self.ir_unit.get() not in ("", "All"): options["unit"] = self.ir_unit.get()
        supplier = getattr(self, "ir_supplier_map", {}).get(self.ir_supplier.get())
        if supplier: options["supplier_id"] = supplier["id"]; options["supplier_name"] = supplier["name"]
        method = {"Weighted average": "average", "FIFO": "fifo"}.get(self.ir_method.get())
        if method: options["method"] = method
        return options

    def run_inventory_report(self):
        try: result = self.client.inventory_report(REPORTS[self.ir_report.get()], self.inventory_report_options())
        except Exception as exc: return messagebox.showerror("Inventory Reports", str(exc))
        self.inventory_report_result = result; self.show_sections(self.ir_viewer, result["sections"]); self.ir_info.config(text=f'{result["title"]}  |  ' + "   ".join(result["meta"]))

    def export_inventory_report(self, format_name):
        if not getattr(self, "inventory_report_result", None): self.run_inventory_report()
        result = getattr(self, "inventory_report_result", None)
        if result: self.output_sections(result["title"], result["meta"], result["sections"], result["title"].replace(" ", "_"), format_name)

    # ------------------------------------------------------------ warehouses & settings
    def build_inventory_settings_page(self, page):
        box = tk.LabelFrame(page, text="Inventory settings", bg=LIGHT, padx=8, pady=6); box.pack(fill="x", padx=8, pady=6)
        self.inv_currency = tk.StringVar(value="USD"); self.inv_method = tk.StringVar(value="Weighted average")
        tk.Label(box, text="Stock valued in", bg=LIGHT).pack(side="left"); ttk.Combobox(box, textvariable=self.inv_currency, values=["USD", "LBP", "EUR", "AED"], state="readonly", width=6).pack(side="left", padx=(4, 10))
        tk.Label(box, text="Costing method", bg=LIGHT).pack(side="left"); ttk.Combobox(box, textvariable=self.inv_method, values=["Weighted average", "FIFO"], state="readonly", width=16).pack(side="left", padx=(4, 10))
        self.action_button(box, "Save Settings", self.save_inventory_settings).pack(side="left", padx=4)
        year = getattr(self, "current_fiscal_year", datetime.now().year)
        tk.Button(box, text=f"Post Stock Variation {year}", command=self.post_stock_variation, bg=GOLD, fg=NAVY, border=0, padx=14, pady=6, font=("Segoe UI", 9, "bold")).pack(side="left", padx=(20, 4))
        tk.Label(page, text="Lebanese periodic method: purchases stay in 601. The Stock Variation voucher (type 06) cancels the stock in account 37 against 6051 and books the counted "
                 "closing stock (Dr 37 / Cr 6052). It is made automatically when you close the year, and the closing stock becomes the Opening Stock of the next year.",
                 bg=LIGHT, fg=MUTED, wraplength=1080, justify="left").pack(fill="x", padx=12, pady=(0, 6))
        wh = tk.LabelFrame(page, text="Warehouses", bg=LIGHT, padx=8, pady=6); wh.pack(fill="both", expand=True, padx=8, pady=6)
        self.wh_id = None; self.wh_code = tk.StringVar(); self.wh_name = tk.StringVar(); self.wh_active = tk.BooleanVar(value=True)
        form = tk.Frame(wh, bg=LIGHT); form.pack(fill="x")
        tk.Label(form, text="Code (blank = automatic)", bg=LIGHT).pack(side="left"); tk.Entry(form, textvariable=self.wh_code, width=8).pack(side="left", padx=(4, 10))
        tk.Label(form, text="Name", bg=LIGHT).pack(side="left"); tk.Entry(form, textvariable=self.wh_name, width=28).pack(side="left", padx=(4, 10))
        tk.Checkbutton(form, text="Active", variable=self.wh_active, bg=LIGHT).pack(side="left")
        self.action_button(form, "New", lambda: (setattr(self, "wh_id", None), self.wh_code.set(""), self.wh_name.set(""), self.wh_active.set(True))).pack(side="left", padx=3)
        tk.Button(form, text="Save Warehouse", command=self.save_warehouse, bg=GOLD, fg=NAVY, border=0, padx=14, pady=6, font=("Segoe UI", 9, "bold")).pack(side="left", padx=3)
        self.warehouses_tree = ttk.Treeview(wh, columns=("code", "name", "active"), show="headings", height=6)
        for key, label, width in (("code", "Code", 90), ("name", "Warehouse", 300), ("active", "Active", 80)): self.warehouses_tree.heading(key, text=label); self.warehouses_tree.column(key, width=width)
        self.warehouses_tree.pack(fill="both", expand=True, pady=6); self.warehouses_tree.bind("<Double-1>", lambda _e: self.edit_warehouse())

    def edit_warehouse(self):
        selected = self.warehouses_tree.selection()
        if not selected: return
        w = next(w for w in self.warehouse_rows if str(w["id"]) == selected[0]); self.wh_id = w["id"]; self.wh_code.set(w["code"]); self.wh_name.set(w["name"]); self.wh_active.set(bool(w["active"]))

    def save_warehouse(self):
        try: self.client.save_warehouse({"id": self.wh_id, "code": self.wh_code.get(), "name": self.wh_name.get(), "active": self.wh_active.get()})
        except Exception as exc: return messagebox.showerror("Warehouses", str(exc))
        self.wh_id = None; self.wh_code.set(""); self.wh_name.set(""); self.load_inventory()

    def save_inventory_settings(self):
        try: self.client.save_inventory_settings({"currency": self.inv_currency.get(), "method": "fifo" if self.inv_method.get() == "FIFO" else "average"})
        except Exception as exc: return messagebox.showerror("Inventory", str(exc))
        self.load_inventory(); messagebox.showinfo("Inventory", "Inventory settings saved")

    def post_stock_variation(self):
        year = getattr(self, "current_fiscal_year", datetime.now().year)
        if not messagebox.askyesno("Stock Variation", f"Post (or replace) the Stock Variation voucher of {year} with the stock value at 31-12-{year}?"): return
        try: result = self.client.post_stock_variation(year)
        except Exception as exc: return messagebox.showerror("Stock Variation", str(exc))
        messagebox.showinfo("Stock Variation", f"Voucher {result.get('voucher') or '-'}\nStock in the ledger before: {result['opening']:,.2f}\nClosing stock: {result['closing']:,.2f}")
        self.load_journal(); self.load_trial()

    # ------------------------------------------------------------ lists for the item form and filters
    def load_item_lists(self):
        try: data = self.client.item_categories(); parties = self.client.parties()
        except Exception: return
        self.category_data = data; suppliers = [p for p in parties if p["kind"] in ("supplier", "both")]
        self.ir_supplier_map = {p["name"]: p for p in suppliers}
        categories = [c["name"] for c in data["categories"]]; subcategories = sorted({s for c in data["categories"] for s in c["subcategories"]})
        if hasattr(self, "item_boxes"):
            self.item_boxes["category"]["values"] = categories; self.item_boxes["unit"]["values"] = data["units"]; self.item_boxes["supplier_name"]["values"] = [p["name"] for p in suppliers]
            self.item_category_chosen()
        for name, values in (("ir_category_box", ["All"] + categories), ("ir_subcategory_box", ["All"] + subcategories), ("ir_unit_box", ["All"] + data["units"]),
                             ("ir_supplier_box", ["All"] + [p["name"] for p in suppliers])):
            box = getattr(self, name, None)
            if box is not None and box.winfo_exists(): box["values"] = values
        if hasattr(self, "cat_tree") and self.cat_tree.winfo_exists():
            self.cat_tree.delete(*self.cat_tree.get_children())
            for c in data["categories"]:
                parent = self.cat_tree.insert("", "end", text=c["name"], values=("Category",), open=True)
                for sub in c["subcategories"]: self.cat_tree.insert(parent, "end", text=sub, values=("Subcategory",))
            self.unit_list.delete(0, "end"); [self.unit_list.insert("end", u) for u in data["units"]]
            self.cat_parent_box["values"] = categories

    def item_category_chosen(self):
        data = getattr(self, "category_data", None)
        if not data or not hasattr(self, "item_boxes"): return
        chosen = next((c for c in data["categories"] if c["name"] == self.item_vars["category"].get()), None)
        self.item_boxes["subcategory"]["values"] = chosen["subcategories"] if chosen else []

    # ------------------------------------------------------------ categories, subcategories, units
    def build_categories_page(self, page):
        left = tk.LabelFrame(page, text="Categories and subcategories", bg=LIGHT, padx=8, pady=6); left.pack(side="left", fill="both", expand=True, padx=8, pady=8)
        self.cat_name = tk.StringVar(); self.cat_parent = tk.StringVar(); self.unit_name = tk.StringVar()
        form = tk.Frame(left, bg=LIGHT); form.pack(fill="x")
        tk.Label(form, text="New category", bg=LIGHT).grid(row=0, column=0, sticky="w"); tk.Entry(form, textvariable=self.cat_name, width=22).grid(row=0, column=1, padx=4)
        self.action_button(form, "Add Category", lambda: self.save_category("category")).grid(row=0, column=2, padx=4)
        self.sub_name = tk.StringVar()
        tk.Label(form, text="New subcategory", bg=LIGHT).grid(row=1, column=0, sticky="w", pady=4); tk.Entry(form, textvariable=self.sub_name, width=22).grid(row=1, column=1, padx=4)
        tk.Label(form, text="in", bg=LIGHT).grid(row=1, column=2, sticky="w"); self.cat_parent_box = ttk.Combobox(form, textvariable=self.cat_parent, state="readonly", width=18); self.cat_parent_box.grid(row=1, column=3, padx=4)
        self.action_button(form, "Add Subcategory", lambda: self.save_category("subcategory")).grid(row=1, column=4, padx=4)
        self.cat_tree = ttk.Treeview(left, columns=("level",), height=14); self.cat_tree.heading("#0", text="Name"); self.cat_tree.heading("level", text="Level")
        self.cat_tree.column("#0", width=260); self.cat_tree.pack(fill="both", expand=True, pady=6)
        right = tk.LabelFrame(page, text="Units", bg=LIGHT, padx=8, pady=6); right.pack(side="left", fill="y", padx=8, pady=8)
        row = tk.Frame(right, bg=LIGHT); row.pack(fill="x")
        tk.Entry(row, textvariable=self.unit_name, width=14).pack(side="left"); self.action_button(row, "Add Unit", lambda: self.save_category("unit")).pack(side="left", padx=4)
        self.unit_list = tk.Listbox(right, height=16, width=22); self.unit_list.pack(fill="y", expand=True, pady=6)

    def save_category(self, kind):
        name = {"category": self.cat_name, "subcategory": self.sub_name, "unit": self.unit_name}[kind].get().strip()
        try: self.client.save_item_category({"kind": kind, "name": name, "parent": self.cat_parent.get()})
        except Exception as exc: return messagebox.showerror("Categories & Units", str(exc))
        {"category": self.cat_name, "subcategory": self.sub_name, "unit": self.unit_name}[kind].set(""); self.load_item_lists()

    # ------------------------------------------------------------ stock in / stock out
    def build_stock_in_out_page(self, page):
        self.sio_vars = {k: tk.StringVar() for k in ("direction", "date", "warehouse", "reason")}
        self.sio_vars["direction"].set("Stock In (add)"); self.sio_vars["date"].set(datetime.now().strftime("%d-%m-%Y"))
        bar = tk.Frame(page, bg=LIGHT); bar.pack(fill="x", padx=8, pady=6)
        ttk.Combobox(bar, textvariable=self.sio_vars["direction"], values=["Stock In (add)", "Stock Out (remove)"], state="readonly", width=18).pack(side="left", padx=(0, 8))
        tk.Label(bar, text="Date", bg=LIGHT).pack(side="left"); self.date_entry(bar, self.sio_vars["date"], 11).pack(side="left", padx=(4, 8))
        tk.Label(bar, text="Warehouse", bg=LIGHT).pack(side="left"); self.sio_wh_box = ttk.Combobox(bar, textvariable=self.sio_vars["warehouse"], state="readonly", width=18); self.sio_wh_box.pack(side="left", padx=(4, 8))
        tk.Label(bar, text="Reason", bg=LIGHT).pack(side="left"); tk.Entry(bar, textvariable=self.sio_vars["reason"], width=28).pack(side="left", padx=4)
        bottom = tk.Frame(page, bg=LIGHT); bottom.pack(side="bottom", fill="x", padx=8, pady=6)
        self.action_button(bottom, "Add Line", lambda: self.sio_sheet.insert(self.sio_row())).pack(side="left", padx=(0, 3))
        tk.Button(bottom, text="Delete Line", command=lambda: self.sio_sheet.delete_selected(), bg=RED, fg="white", border=0, padx=10, pady=6).pack(side="left", padx=3)
        tk.Button(bottom, text="Save", command=self.save_stock_in_out, bg=GOLD, fg=NAVY, border=0, padx=18, pady=6, font=("Segoe UI", 9, "bold")).pack(side="left", padx=(12, 3))
        tk.Label(bottom, text="Stock In: cost = average cost unless you type another. Stock Out: leaves at the average cost. Quantities and costs update at once.", bg=LIGHT, fg=MUTED).pack(side="left", padx=8)
        self.sio_sheet = EditableSheet(self, page, [("line", "#", 40, "center"), ("sku", "Item Code", 110, "w"), ("name", "Item", 260, "w"), ("unit", "Unit", 60, "center"),
            ("on_hand", "On Hand", 100, "e"), ("quantity", "Quantity", 100, "e"), ("unit_cost", "Unit Cost (average)", 130, "e"), ("value", "Value", 120, "e")],
            ["sku", "quantity", "unit_cost"], self.sio_changed, height=10)
        self.sio_sheet.tree.bind("<F2>", lambda _e: self.item_picker(lambda sku: self.sio_pick(sku)))
        for _ in range(3): self.sio_sheet.insert(self.sio_row())

    def sio_row(self):
        return {"sku": "", "name": "", "unit": "", "on_hand": "", "quantity": "", "unit_cost": "", "value": ""}

    def sio_pick(self, sku):
        iid, _row = self.sio_sheet.selected()
        if iid: self.sio_changed(iid, "sku", sku); self.sio_sheet.refresh(iid)

    def sio_changed(self, iid, key, text):
        row = self.sio_sheet.rows[iid]
        if key == "sku":
            item = self.item_by_code(text)
            if text and not item: messagebox.showwarning("Stock In / Out", f"Item {text} was not found (F2 to search)"); return False
            if item: row.update(sku=item["sku"], name=item["name"], unit=item["unit"], on_hand=f'{item["quantity"]:,.3f}', unit_cost=round(item["average_cost"], 4))
        else:
            value = _num(text)
            if value is None or value < 0: messagebox.showwarning("Stock In / Out", "Enter a positive number"); return False
            if key == "unit_cost" and self.sio_vars["direction"].get().startswith("Stock Out"): messagebox.showinfo("Stock In / Out", "Stock out leaves at the average cost"); return False
            row[key] = value
        row["value"] = f'{(_num(row.get("quantity")) or 0) * (_num(row.get("unit_cost")) or 0):,.2f}'

    def save_stock_in_out(self):
        lines = [r for r in self.sio_sheet.ordered() if r.get("sku") and _num(r.get("quantity"))]
        if not lines: return messagebox.showwarning("Stock In / Out", "Add at least one item with a quantity")
        stock_in = self.sio_vars["direction"].get().startswith("Stock In")
        header = {"doc_type": "adjustment_in" if stock_in else "adjustment_out", "doc_date": self.sio_vars["date"].get(), "warehouse_id": (self.sio_vars["warehouse"].get() or "MAIN").split(" - ", 1)[0],
                  "notes": self.sio_vars["reason"].get().strip() or ("Stock in" if stock_in else "Stock out")}
        try: saved = self.client.save_stock_document(header, [{"sku": r["sku"], "quantity": r["quantity"], "unit_cost": r.get("unit_cost") if stock_in else 0} for r in lines])
        except Exception as exc: return messagebox.showerror("Stock In / Out", str(exc))
        messagebox.showinfo("Stock In / Out", f'{saved["number"]} saved'); self.sio_sheet.clear(); [self.sio_sheet.insert(self.sio_row()) for _ in range(3)]; self.load_inventory()

    def item_picker(self, callback):
        window = tk.Toplevel(self); window.title("Items - F2"); window.geometry("660x420"); window.configure(bg=LIGHT); window.transient(self); window.grab_set()
        search = tk.StringVar(); entry = tk.Entry(window, textvariable=search, width=40); entry.pack(padx=10, pady=8); entry.focus_set()
        tree = ttk.Treeview(window, columns=("sku", "name", "category", "qty"), show="headings"); tree.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        for key, label, width in (("sku", "Code", 100), ("name", "Item", 280), ("category", "Category", 120), ("qty", "On Hand", 90)): tree.heading(key, text=label); tree.column(key, width=width)
        def fill(*_a):
            tree.delete(*tree.get_children()); text = search.get().casefold()
            for i in getattr(self, "inventory_rows", []):
                if i["active"] and (not text or text in f'{i["sku"]} {i["name"]} {i.get("category") or ""} {i.get("barcode") or ""}'.casefold()): tree.insert("", "end", values=(i["sku"], i["name"], i.get("category") or "", f'{i["quantity"]:,.3f}'))
        def choose(_e=None):
            if tree.selection(): value = tree.item(tree.selection()[0], "values")[0]; window.destroy(); callback(value)
        search.trace_add("write", fill); tree.bind("<Double-1>", choose); tree.bind("<Return>", choose); entry.bind("<Return>", lambda _e: (tree.selection_set(tree.get_children()[0]), choose()) if tree.get_children() else None); fill()

    # ------------------------------------------------------------ physical inventory
    def build_physical_page(self, page):
        self.pc_id = None; self.pc_vars = {k: tk.StringVar() for k in ("date", "warehouse", "find")}; self.pc_vars["date"].set(datetime.now().strftime("%d-%m-%Y"))
        bar = tk.Frame(page, bg=LIGHT); bar.pack(fill="x", padx=8, pady=6)
        tk.Label(bar, text="Count date", bg=LIGHT).pack(side="left"); self.date_entry(bar, self.pc_vars["date"], 11).pack(side="left", padx=(4, 8))
        tk.Label(bar, text="Warehouse", bg=LIGHT).pack(side="left"); self.pc_wh_box = ttk.Combobox(bar, textvariable=self.pc_vars["warehouse"], state="readonly", width=18); self.pc_wh_box.pack(side="left", padx=(4, 8))
        tk.Button(bar, text="Load Stock on Hand", command=self.load_count_sheet, bg=GOLD, fg=NAVY, border=0, padx=12, pady=5, font=("Segoe UI", 9, "bold")).pack(side="left", padx=3)
        tk.Label(bar, text="Open count", bg=LIGHT).pack(side="left", padx=(12, 2)); self.pc_find_box = ttk.Combobox(bar, textvariable=self.pc_vars["find"], state="readonly", width=34); self.pc_find_box.pack(side="left")
        self.pc_find_box.bind("<<ComboboxSelected>>", lambda _e: self.open_count())
        bottom = tk.Frame(page, bg=LIGHT); bottom.pack(side="bottom", fill="x", padx=8, pady=6)
        self.action_button(bottom, "Save Count", lambda: self.save_count(False)).pack(side="left", padx=(0, 3))
        tk.Button(bottom, text="Post Differences to Stock", command=lambda: self.save_count(True), bg=GOLD, fg=NAVY, border=0, padx=14, pady=6, font=("Segoe UI", 9, "bold")).pack(side="left", padx=3)
        self.action_button(bottom, "Print / PDF", lambda: self.count_document("preview")).pack(side="left", padx=3)
        self.action_button(bottom, "Excel Count Sheet", self.export_count_excel).pack(side="left", padx=3)
        self.action_button(bottom, "Upload Counts (Excel)", self.upload_count_excel).pack(side="left", padx=3)
        self.pc_info = tk.Label(bottom, text="", bg=LIGHT, fg=NAVY, font=("Segoe UI", 9, "bold")); self.pc_info.pack(side="right", padx=6)
        self.pc_sheet = EditableSheet(self, page, [("line", "#", 40, "center"), ("sku", "Item Code", 100, "w"), ("name", "Item", 240, "w"), ("location", "Location", 80, "center"),
            ("unit", "Unit", 55, "center"), ("system_qty", "Stock on Hand", 110, "e"), ("counted", "Physical Count", 110, "e"), ("difference", "Difference", 100, "e"),
            ("value", "Difference Value", 120, "e")], ["counted"], self.count_changed, height=12)

    def load_count_sheet(self):
        warehouse = self.warehouse_id_of(self.pc_vars["warehouse"].get())
        try: rows = self.client.count_sheet(warehouse, self.pc_vars["date"].get())
        except Exception as exc: return messagebox.showerror("Physical Inventory", str(exc))
        self.pc_id = None; self.pc_sheet.clear()
        for row in rows: row["counted"] = ""; self.count_display(row); self.pc_sheet.insert(row)
        self.update_count_info()

    def warehouse_id_of(self, label):
        code = (label or "MAIN").split(" - ", 1)[0]
        return next((w["id"] for w in getattr(self, "warehouse_rows", []) if w["code"] == code), 1)

    def count_display(self, row):
        counted = _num(row.get("counted")) if row.get("counted") not in ("", None) else None
        difference = (counted - row["system_qty"]) if counted is not None else None
        row["difference"] = difference; row["_display"] = {"system_qty": f'{row["system_qty"]:,.3f}', "counted": f"{counted:,.3f}" if counted is not None else "",
            "difference": f"{difference:+,.3f}" if difference else ("0" if difference == 0 else ""), "value": f'{difference * row["unit_cost"]:+,.2f}' if difference else ""}

    def count_changed(self, iid, key, text):
        row = self.pc_sheet.rows[iid]
        if text.strip() == "": row["counted"] = ""
        else:
            value = _num(text)
            if value is None or value < 0: messagebox.showwarning("Physical Inventory", "Enter the counted quantity (0 or more)"); return False
            row["counted"] = value
        self.count_display(row); self.update_count_info()

    def update_count_info(self):
        rows = self.pc_sheet.ordered(); counted = [r for r in rows if r.get("counted") not in ("", None)]
        value = sum((r["difference"] or 0) * r["unit_cost"] for r in counted)
        self.pc_info.config(text=f"{len(counted)} of {len(rows)} item(s) counted   Difference value: {value:+,.2f} {getattr(self, 'inventory_currency', '')}")

    def save_count(self, post):
        lines = [{"item_id": r["item_id"], "sku": r["sku"], "counted": r["counted"]} for r in self.pc_sheet.ordered() if r.get("counted") not in ("", None)]
        if not lines: return messagebox.showwarning("Physical Inventory", "Enter at least one physical count")
        if post and not messagebox.askyesno("Physical Inventory", "Post the differences? Adjustment + / - documents are made so the stock equals the physical count."): return
        try: saved = self.client.save_physical_count({"count_date": self.pc_vars["date"].get(), "warehouse_id": self.warehouse_id_of(self.pc_vars["warehouse"].get())}, lines, self.pc_id, post)
        except Exception as exc: return messagebox.showerror("Physical Inventory", str(exc))
        self.pc_id = saved["id"]
        messagebox.showinfo("Physical Inventory", f'{saved["number"]} saved' + (f' and posted: {saved.get("adjustment_numbers") or "no difference"}' if post else " (draft)"))
        self.load_inventory(); self.load_counts()

    def load_counts(self):
        try: counts = self.client.physical_counts()
        except Exception: counts = []
        self.pc_map = {f'{c["number"]} | {_dd(c["count_date"])} | {c["warehouse_code"]} | {c["status"]}': c["id"] for c in counts}; self.pc_find_box["values"] = list(self.pc_map)

    def open_count(self):
        count_id = getattr(self, "pc_map", {}).get(self.pc_vars["find"].get())
        if not count_id: return
        try: count = self.client.physical_count(count_id); rows = self.client.count_sheet(count["warehouse_id"], count["count_date"])
        except Exception as exc: return messagebox.showerror("Physical Inventory", str(exc))
        counted = {l["item_id"]: l["counted"] for l in count["lines"]}; self.pc_id = None if count["status"] == "posted" else count["id"]; self.pc_sheet.clear()
        self.pc_vars["date"].set(_dd(count["count_date"]))
        for row in rows: row["counted"] = float(counted[row["item_id"]]) if row["item_id"] in counted else ""; self.count_display(row); self.pc_sheet.insert(row)
        self.update_count_info()
        if count["status"] == "posted": self.pc_info.config(text=f'{count["number"]} is posted ({count.get("adjustment_numbers") or "no difference"}): shown for printing')

    def count_rows_for_export(self):
        return [[r["sku"], r["name"], r.get("location") or "", r["unit"], r["system_qty"], r["counted"] if r.get("counted") not in ("", None) else "", r["difference"] if r.get("difference") is not None else "",
                 round((r["difference"] or 0) * r["unit_cost"], 2) if r.get("difference") else ""] for r in self.pc_sheet.ordered()]

    def count_document(self, mode):
        rows = self.count_rows_for_export()
        if not rows: return messagebox.showwarning("Physical Inventory", "Load the stock on hand first")
        self.output_sections("Physical Inventory Count", [f"Date: {self.pc_vars['date'].get()}   Warehouse: {self.pc_vars['warehouse'].get()}", "Counted by: ______________   Checked by: ______________"],
            [{"heading": "Count sheet", "headers": ["Item Code", "Item", "Location", "Unit", "Stock on Hand", "Physical Count", "Difference", "Difference Value"], "rows": rows, "total_rows": []}], "Physical_Count", mode)

    def export_count_excel(self):
        from report_export import export_excel
        rows = self.count_rows_for_export()
        if not rows: return messagebox.showwarning("Physical Inventory", "Load the stock on hand first")
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile="Physical_Count.xlsx", filetypes=[("Excel", "*.xlsx")])
        if not path: return
        export_excel(path, "Physical Count", ["Item Code", "Item", "Location", "Unit", "Stock on Hand", "Physical Count", "Difference", "Difference Value"], rows)
        messagebox.showinfo("Physical Inventory", f"Saved: {path}\nFill the 'Physical Count' column and use 'Upload Counts (Excel)'.")

    def upload_count_excel(self):
        from openpyxl import load_workbook
        path = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx *.xlsm")])
        if not path: return
        if not self.pc_sheet.ordered(): self.load_count_sheet()
        workbook = load_workbook(path, read_only=True, data_only=True); found = 0
        try:
            rows = list(workbook.worksheets[0].iter_rows(values_only=True))
            header_index = next(i for i, r in enumerate(rows[:10]) if r and any(str(v or "").strip().lower() in ("item code", "code", "sku") for v in r))
            header = [str(v or "").strip().lower() for v in rows[header_index]]
            code_col = next(i for i, h in enumerate(header) if h in ("item code", "code", "sku")); count_col = next(i for i, h in enumerate(header) if "count" in h or h in ("qty", "quantity"))
            counts = {str(r[code_col]).strip().upper(): r[count_col] for r in rows[header_index + 1:] if r and r[code_col] not in (None, "") and r[count_col] not in (None, "")}
        except StopIteration: return messagebox.showerror("Physical Inventory", "The Excel file needs an 'Item Code' column and a 'Physical Count' column")
        finally: workbook.close()
        for iid, row in self.pc_sheet.rows.items():
            if row["sku"].upper() in counts:
                try: row["counted"] = float(str(counts[row["sku"].upper()]).replace(",", "")); found += 1
                except ValueError: continue
                self.count_display(row); self.pc_sheet.refresh(iid)
        self.update_count_info(); messagebox.showinfo("Physical Inventory", f"{found} count(s) loaded from the Excel file. Check them, then Save or Post.")

